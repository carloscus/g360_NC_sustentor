from pathlib import Path
from typing import Optional
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime
from src.domain import RecognitionResult, ExpedienteComercial
from src.core.utils import coerce_str
from src.core.fechas import fecha_ui
from src.ui.reconocimiento_config import NAMING
from src.ui.catalog import CATALOGO
from src.render.g360_styles import (
    EXP_NAVY,
    EXP_GRAY,
    EXP_FAINT,
    EXP_LABEL_BG,
    EXP_INK,
    EXP_BORDER,
    exp_rgb,
)

NAVY = exp_rgb(EXP_NAVY)
GRAY = exp_rgb(EXP_GRAY)


def _set_cell(p, text, bold=False, size=11, color=None):
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    return run


def _recortar(texto: str, limite: int = 800) -> str:
    """Recorta texto libre al límite para sostener el informe en una hoja."""
    t = str(texto or "").strip()
    if len(t) > limite:
        return t[:limite].rstrip() + "…"
    return t


def _normalizar_texto(texto: str) -> str:
    """Normalización ligera de texto libre.

    Colapsa espacios, corrige errores frecuentes del dominio (tildes y el
    typo 'prercios'), capitaliza el inicio y cierra con punto. No reescribe
    la redacción del usuario.
    """
    t = " ".join(str(texto or "").split())
    if not t:
        return ""
    for patron, fijo in _CORRECCIONES:
        t = re.sub(patron, fijo, t)
    t = t[0].upper() + t[1:]
    if t[-1] not in ".…;:!?":
        t += "."
    return t


_CORRECCIONES = [
    (r"\bprercios\b", "precios"),
    (r"\bPrercios\b", "Precios"),
    (r"\bPRERCIOS\b", "PRECIOS"),
    (r"\batencion\b", "atención"),
    (r"\bAtencion\b", "Atención"),
    (r"\bATENCION\b", "ATENCIÓN"),
    (r"\bsegun\b", "según"),
    (r"\bSegun\b", "Según"),
    (r"\bSEGUN\b", "SEGÚN"),
    (r"\bomision\b", "omisión"),
    (r"\bOmision\b", "Omisión"),
    (r"\bOMISION\b", "OMISIÓN"),
    (r"\bomisiones\b", "omisiones"),
    (r"\bOmisiones\b", "Omisiones"),
    (r"\bOMISIONES\b", "OMISIONES"),
]


def _lista_docs(df) -> list:
    """Facturas únicas sin repetir (de FACTURAS y FACTURA), ordenadas."""
    docs = set()
    if df is not None and not df.empty:
        for col in ("FACTURAS", "FACTURA"):
            if col not in df.columns:
                continue
            for v in df[col].dropna():
                for p in str(v).replace(";", ",").split(","):
                    p = p.strip().split(" (")[0].strip()
                    if p and p.lower() not in ("nan", "none"):
                        docs.add(p)
    return sorted(docs)


def _texto_facturas(df, limite: int = 110) -> str:
    """Facturas comprometidas sin repetir, en una línea (con truncado)."""
    docs = _lista_docs(df)
    if not docs:
        return "—"
    txt = ", ".join(docs)
    if len(txt) <= limite:
        return txt
    visibles, resto = [], 0
    for d in docs:
        candidato = ", ".join(visibles + [d])
        if len(candidato) > limite and visibles:
            resto = len(docs) - len(visibles)
            break
        visibles.append(d)
    txt = ", ".join(visibles)
    return f"{txt} (+{resto} más)" if resto else txt


class DocxInformeRenderer:
    """Genera el Informe de Sustento Comercial como documento programatico.
    Construido desde cero con python-docx."""

    JUSTIFICACION_POR_TIPO = {
        k: v["justificacion"] for k, v in NAMING.items() if "justificacion" in v
    }
    LABELS_TIPO = {k: v["titulo"] for k, v in NAMING.items()}

    def __init__(self):
        pass

    @staticmethod
    def _extraer_ruc(df) -> str:
        if df.empty or "DOC_CLIENTE" not in df.columns:
            return ""
        for val in df["DOC_CLIENTE"].dropna().unique():
            s = str(val).strip()
            if s and s.lower() not in ("nan", "none", ""):
                return s
        return ""

    @staticmethod
    def _build_alcance(df) -> str:
        """Alcance sin montos: '3 SKU · 8,000 unidades' (el dinero va en §2)."""
        if df is None or df.empty:
            return "—"
        monto_cols = ["MONTO_NC", "Subtotal NC (S/)", "MONTO_FACTURA"]
        monto_col = next((c for c in monto_cols if c in df.columns), None)
        base = df[df[monto_col] > 0] if monto_col else df
        n_skus = len(base["SKU"].unique()) if "SKU" in base.columns else len(base)
        total_u = int(base["CANTIDAD"].sum()) if "CANTIDAD" in base.columns else 0
        parts = [f"{n_skus} SKU"]
        if total_u:
            parts.append(f"{total_u:,} unidades")
        return " · ".join(parts)

    def _add_shading(self, element, color_hex):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), color_hex)
        shd.set(qn("w:val"), "clear")
        element.append(shd)

    def _add_bottom_border(self, pPr, color_hex=EXP_NAVY, sz="8"):
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), sz)
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), color_hex)
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _estilo_tabla_apa(self, table):
        """Bordes estilo APA 7.ª ed.: solo líneas horizontales (título de
        tabla arriba, cierre de encabezado y cierre de tabla)."""
        tblPr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            if edge in ("top", "bottom"):
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "6")
                el.set(qn("w:color"), EXP_INK)
            elif edge == "insideH":
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), EXP_BORDER)
            else:
                el.set(qn("w:val"), "nil")
                el.set(qn("w:sz"), "0")
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "auto")
            borders.append(el)
        tblPr.append(borders)

    def _configurar_encabezado(self, doc):
        header = doc.sections[0].header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p._p.get_or_add_pPr()
        self._add_shading(pPr, EXP_NAVY)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "80")
        spacing.set(qn("w:after"), "80")
        pPr.append(spacing)
        run = p.add_run("SUSTENTO COMERCIAL")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = "Calibri"

    def _configurar_pie(self, doc, doc_ref=""):
        footer = doc.sections[0].footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = p._p.get_or_add_pPr()
        self._add_bottom_border(pPr, sz="4")
        label = f"Ref. {doc_ref}" if doc_ref else "G360"
        run = p.add_run(f"{label}  |  Pág. ")
        run.font.size = Pt(8)
        run.font.color.rgb = GRAY
        run.font.name = "Calibri"
        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")
        run._r.append(fldChar1)
        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = " PAGE "
        run._r.append(instrText)
        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")
        run._r.append(fldChar2)
        run2 = p.add_run(f"  |  {fecha_ui(datetime.now())}")
        run2.font.size = Pt(8)
        run2.font.color.rgb = GRAY
        run2.font.name = "Calibri"

    def generar(
        self,
        resultado: RecognitionResult,
        expediente: Optional[ExpedienteComercial] = None,
        ruta_salida: str = "",
        datos_adicionales: Optional[dict] = None,
        nombre_archivo_excel: str = "",
    ) -> Path:
        datos = datos_adicionales or {}
        doc = Document()
        doc.core_properties.author = "ccusi"
        doc.core_properties.description = "Generado por G360"

        # ── Hoja única A4: márgenes estrechos ────────────────────────
        sec = doc.sections[0]
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(1.4)
        sec.bottom_margin = Cm(1.4)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
        sec.header_distance = Cm(0.8)
        sec.footer_distance = Cm(0.8)

        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2)

        self._configurar_encabezado(doc)
        self._configurar_pie(doc, datos.get("numero_referencia", ""))

        if not ruta_salida:
            ruta_salida = f"Informe_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"

        out_path = Path(ruta_salida)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        total_nc = resultado.resumen.get("total_nc", 0)
        doc_ref = resultado.resumen.get("doc_ref", "") or resultado.resumen.get(
            "documento_referencia", ""
        )
        tipo = datos.get("tipo_operacion", "")
        cliente = datos.get("cliente", "")

        if expediente:
            df_datos = expediente.datos
        else:
            df_datos = resultado.dataframe

        df = resultado.dataframe

        # Total consistente con la hoja Cálculo: se recalcula desde el mismo
        # dataframe y el mismo redondeo que usa ExcelRenderer, para que el
        # RESUMEN ECONÓMICO nunca difiera del subtotal de la hoja.
        if not df.empty:
            monto_col = next(
                (
                    c
                    for c in ("MONTO_NC", "Subtotal NC (S/)", "SUBTOTAL (SIN IGV)")
                    if c in df.columns
                ),
                None,
            )
            if monto_col:
                try:
                    total_nc = float(df[monto_col].apply(lambda x: round(float(x), 2)).sum())
                except (TypeError, ValueError):
                    pass

        ruc = self._extraer_ruc(df_datos)
        igv = round(total_nc * 0.18, 2)
        total_con_igv = round(total_nc + igv, 2)
        descripcion = datos.get("descripcion", "")
        observaciones = datos.get("observaciones", "")
        vendedor = datos.get("representante", "")

        self._agregar_titulo(doc, doc_ref, tipo)
        modalidad = datos.get("modalidad", "individual")
        # Una sola fila de referencia (sin redundancia): el documento en
        # individual, la lista sin repetir en consolidado.
        ref_txt = doc_ref
        if modalidad == "consolidado" or not ref_txt:
            ref_txt = _texto_facturas(df)
        self._agregar_datos_generales(
            doc,
            cliente,
            doc_ref,
            vendedor,
            ruc,
            tipo,
            periodo=datos.get("periodo", ""),
            alcance_txt=self._build_alcance(df),
            ref_txt=ref_txt,
            fecha_doc=datos.get("fecha_documento", ""),
        )
        tipo_leg = self._tipo_legacy(datos.get("tipo_calculo", "") or tipo)
        self._agregar_resultado(doc, total_nc, igv, total_con_igv)
        frase = self._frase_cierre(tipo_leg, modalidad, df, total_nc)
        self._agregar_analisis(doc, tipo, descripcion, frase)
        self._agregar_observaciones(doc, observaciones)
        self._agregar_documentos(
            doc, nombre_archivo_excel, Path(out_path).name, datos.get("nombre_historico", "")
        )

        doc.save(str(out_path))
        return out_path

    def _agregar_titulo(self, doc: Document, doc_ref: str = "", tipo: str = ""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(0)
        base_titulo = NAMING.get(tipo, {}).get("docx_titulo", "INFORME DE SUSTENTO COMERCIAL")
        if doc_ref:
            base_titulo += f"  |  {doc_ref}"
        run = p.add_run(base_titulo)
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = "Calibri"
        run.font.color.rgb = NAVY
        subtitulo = NAMING.get(tipo, {}).get("subtitulo", "")
        if subtitulo:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(2)
            run2 = p2.add_run(subtitulo)
            run2.font.size = Pt(9)
            run2.font.name = "Calibri"
            run2.font.color.rgb = exp_rgb(EXP_FAINT)

    def _agregar_datos_generales(
        self,
        doc: Document,
        cliente: str,
        doc_ref: str,
        vendedor: str,
        ruc: str = "",
        tipo: str = "",
        periodo: str = "",
        alcance_txt: str = "",
        ref_txt: str = "",
        fecha_doc: str = "",
    ):
        self._agregar_seccion(doc, "1. DATOS GENERALES")
        table = doc.add_table(rows=8, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.style = "Table Grid"
        self._estilo_tabla_apa(table)

        labels_tipo = self.LABELS_TIPO
        _caso = CATALOGO.get(tipo)
        tipo_label = coerce_str(
            (_caso.label if _caso else labels_tipo.get(tipo, tipo)), "No especificado"
        )

        campos = [
            ("Cliente", cliente or ""),
            ("RUC", ruc or ""),
            ("Vendedor", vendedor or ""),
            ("Periodo", periodo or "—"),
            ("Doc. de referencia", ref_txt or doc_ref or "—"),
            ("Alcance", alcance_txt or "—"),
            ("Tipo de gestión", tipo_label),
            (
                "Fecha de documento" if fecha_doc else "Fecha de elaboración",
                fecha_doc or fecha_ui(datetime.now()),
            ),
        ]
        for i, (campo, valor) in enumerate(campos):
            row = table.rows[i]
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            row.height = Cm(0.55)
            c_lbl, c_val = row.cells[0], row.cells[1]
            c_lbl.text = ""
            c_val.text = ""
            c_lbl.width = Cm(4.2)
            self._add_shading(c_lbl._tc.get_or_add_tcPr(), EXP_LABEL_BG)
            _set_cell(c_lbl.paragraphs[0], campo, bold=True, size=10)
            _set_cell(c_val.paragraphs[0], valor, size=10)

    @staticmethod
    def _tipo_legacy(tipo: str) -> str:
        """Normaliza código de caso ('DC') o legacy a clave de NAMING."""
        if not tipo:
            return ""
        if tipo in NAMING:
            return tipo
        caso = CATALOGO.get(tipo)
        if caso is not None and getattr(caso, "strategy", ""):
            return caso.strategy
        return tipo

    def _agregar_resultado(self, doc, subtotal, igv, total):
        """§2 la conclusión en dinero (el detalle vive en el Cálculo)."""
        self._agregar_seccion(doc, "2. RESULTADO ECONÓMICO")
        total_table = doc.add_table(rows=3, cols=2)
        total_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        total_table.style = "Table Grid"
        self._estilo_tabla_apa(total_table)
        data = [
            ("Subtotal (Sin IGV)", f"S/ {subtotal:,.2f}"),
            ("IGV (18%)", f"S/ {igv:,.2f}"),
            ("Total (con IGV)", f"S/ {total:,.2f}"),
        ]
        for i, (campo, valor) in enumerate(data):
            row = total_table.rows[i]
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            row.height = Cm(0.55)
            row.cells[0].text = ""
            row.cells[1].text = ""
            row.cells[0].width = Cm(4.2)
            self._add_shading(row.cells[0]._tc.get_or_add_tcPr(), EXP_LABEL_BG)
            _set_cell(row.cells[0].paragraphs[0], campo, bold=True, size=10)
            _set_cell(
                row.cells[1].paragraphs[0],
                valor,
                bold=(i == 2),
                size=11 if i == 2 else 10,
                color=NAVY if i == 2 else None,
            )

    @staticmethod
    def _n_skus_con_nc(df) -> int:
        if df is None or df.empty or "SKU" not in df.columns:
            return 0
        monto_col = next(
            (
                c
                for c in ("MONTO_NC", "Subtotal NC (S/)", "SUBTOTAL (SIN IGV)", "MONTO_FACTURA")
                if c in df.columns
            ),
            None,
        )
        base = df[df[monto_col] > 0] if monto_col else df
        return len(base["SKU"].unique())

    def _frase_cierre(self, tipo_leg: str, modalidad: str, df, total: float) -> str:
        """Frase de cierre del análisis, personalizada por tipo de proceso.

        Semántica del negocio: regularización de documentos o promociones —
        la nota de crédito regulariza lo cobrado en exceso o lo omitido.
        """
        skus = self._n_skus_con_nc(df)
        docs = _lista_docs(df)
        total_txt = f"S/ {total:,.2f}"
        if tipo_leg == "diferencia_precio":
            if modalidad == "consolidado" or len(docs) != 1:
                return (
                    f"Se revisó el histórico frente a la lista de precios vigente; "
                    f"la nota de crédito por {total_txt} regulariza las diferencias "
                    f"en {skus} SKU de {len(docs)} factura(s)."
                )
            return (
                f"Se revisó el histórico frente a la lista de precios vigente; "
                f"la nota de crédito por {total_txt} regulariza las diferencias "
                f"en {skus} SKU de la factura {docs[0]}."
            )
        if tipo_leg in ("descuento_precio", "descuento_factura"):
            return (
                f"Se revisaron los descuentos omitidos frente al histórico; "
                f"la nota de crédito por {total_txt} regulariza lo no aplicado "
                f"en {skus} SKU."
            )
        if tipo_leg == "anular_factura":
            docs = _lista_docs(df)
            doc_txt = docs[0] if len(docs) == 1 else f"{len(docs)} factura(s)"
            return (
                f"La factura {doc_txt} se anula al 100%; "
                f"la nota de crédito por {total_txt} regulariza el documento."
            )
        if tipo_leg == "rebate_volumen":
            lineas = (
                sorted({str(v).strip() for v in df["LINEA"].dropna() if str(v).strip()})
                if df is not None and not df.empty and "LINEA" in df.columns
                else []
            )
            if len(lineas) == 1:
                lin_txt = f" en línea {lineas[0]}"
            elif lineas:
                lin_txt = f" en {len(lineas)} líneas"
            else:
                lin_txt = ""
            return (
                f"Se verificó el volumen frente a la meta acordada; "
                f"la nota de crédito por {total_txt} regulariza el rebate{lin_txt}."
            )
        if tipo_leg == "bonificacion_promocion":
            return (
                f"Se verificó la mecánica promocional frente al histórico; "
                f"la nota de crédito por {total_txt} regulariza la bonificación "
                f"en {skus} SKU."
            )
        if tipo_leg == "feria_preventa":
            return (
                f"Se asignaron facturas de sustento al compromiso de feria/preventa; "
                f"la nota de crédito por {total_txt} regulariza lo comprometido "
                f"en {skus} SKU."
            )
        just = self.JUSTIFICACION_POR_TIPO.get(tipo_leg, "")
        base = just or (
            "Se sustenta el reconocimiento comercial mediante la emisión "
            "de la Nota de Crédito por el importe calculado."
        )
        return f"{base} Importe regularizado: {total_txt} en {skus} SKU."

    def _agregar_analisis(self, doc, tipo, antecedentes, frase_cierre: str = ""):
        """§3 narrativa del caso + frase de cierre (sin datos del Cálculo)."""
        self._agregar_seccion(doc, "3. ANTECEDENTES / ANÁLISIS COMERCIAL")

        if antecedentes:
            p = doc.add_paragraph()
            run = p.add_run(_normalizar_texto(_recortar(antecedentes)))
            run.font.size = Pt(10)
            run.font.name = "Calibri"
        else:
            # Fallback to generic justification if no antecedentes provided
            just = self.JUSTIFICACION_POR_TIPO.get(
                tipo,
                (
                    "Se sustenta el reconocimiento comercial mediante la emisión "
                    "de la Nota de Crédito por el importe calculado."
                ),
            )
            p = doc.add_paragraph()
            run = p.add_run(just)
            run.font.size = Pt(10)
            run.font.name = "Calibri"

        if frase_cierre:
            p = doc.add_paragraph()
            run = p.add_run(frase_cierre)
            run.font.size = Pt(10)
            run.font.name = "Calibri"

    def _agregar_observaciones(self, doc, observaciones):
        self._agregar_seccion(doc, "4. OBSERVACIONES")
        p = doc.add_paragraph()
        run = p.add_run(
            _normalizar_texto(_recortar(observaciones)) or "Sin observaciones registradas."
        )
        run.font.size = Pt(10)
        run.font.name = "Calibri"

    def _agregar_documentos(
        self, doc: Document, nombre_excel: str, nombre_informe: str, nombre_historico: str = ""
    ):
        self._agregar_seccion(doc, "5. DOCUMENTOS")
        archivos = [
            (nombre_excel, "Hoja de cálculo"),
            (nombre_historico, "Segmento histórico del ERP"),
            (nombre_informe, "Este informe"),
        ]
        for nombre, rol in archivos:
            if not nombre:
                continue
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(f"{nombre} — {rol}")
            run.font.size = Pt(10)
            run.font.name = "Calibri"

    def _agregar_seccion(self, doc: Document, titulo: str):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        pPr = p._p.get_or_add_pPr()
        self._add_bottom_border(pPr)
        run = p.add_run(titulo)
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = "Calibri"
        run.font.color.rgb = NAVY
