from openpyxl.styles import Font, PatternFill, Border, Side, Alignment


class G360Styles:
    """
    Centraliza la identidad visual de los reportes G360.
    Define colores, fuentes y bordes compartidos entre plantillas y reportes finales.
    """

    def __init__(self):
        self.side = Side(style="thin", color="000000")
        self.border = Border(left=self.side, right=self.side, top=self.side, bottom=self.side)
        self.header_fill = PatternFill(start_color="0D2B4E", end_color="0D2B4E", fill_type="solid")
        self.header_font = Font(color="FFFFFF", bold=True, size=10)
        self.critical_fill = PatternFill(
            start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
        )
        self.total_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
        self.alert_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        self.warning_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        self.info_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        self.zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        self.alert_font = Font(color="9C0006", bold=True, size=10)
        self.warning_font = Font(color="9C5700", bold=True, size=10)
        self.info_font = Font(color="003366", bold=True, size=10)
        self.center_align = Alignment(horizontal="center", vertical="center")
        self.left_align = Alignment(horizontal="left", vertical="center")
        self.right_align = Alignment(horizontal="right", vertical="center")
        self.wrap_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        self.title_font = Font(bold=True, size=14)
        self.title_font_green = Font(bold=True, size=16, color="008554")
        self.bold_font = Font(bold=True)
        self.italic_gray_font = Font(italic=True, color="666666")
        self.trend_up_font = Font(color="008000", bold=True)
        self.trend_down_font = Font(color="FF0000", bold=True)
        self.note_fill = PatternFill(start_color="FFFBE6", end_color="FFFBE6", fill_type="solid")

        self.hhi_high_fill = PatternFill(
            start_color="9C0006", end_color="9C0006", fill_type="solid"
        )
        self.hhi_high_font = Font(bold=True, size=12, color="FFFFFF")
        self.hhi_mod_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        self.hhi_mod_font = Font(bold=True, size=12, color="734000")
        self.hhi_low_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        self.hhi_low_font = Font(bold=True, size=12, color="006100")

        self.subtotal_font = Font(bold=True, size=11)
        self.subtotal_fill = PatternFill(
            start_color="D6E4F0", end_color="D6E4F0", fill_type="solid"
        )
        self.kpi_value_fill = PatternFill(
            start_color="E8F0FE", end_color="E8F0FE", fill_type="solid"
        )
        self.kpi_header_fill = PatternFill(
            start_color="0D2B4E", end_color="0D2B4E", fill_type="solid"
        )
        self.kpi_label_font = Font(color="FFFFFF", bold=True, size=11)
        self.kpi_value_font = Font(bold=True, size=12, color="003366")

        self.dev_font = Font(color="9C0006")
        self.dev_sub_font = Font(bold=True, size=11, color="9C0006")
        self.dev_label_font = Font(bold=True, size=10, color="9C0006")

        self.vuln_high_font = Font(bold=True, color="9C0006")
        self.vuln_mod_font = Font(bold=True, color="9C5700")

        self.critical_white_font = Font(bold=True, color="FFFFFF")

        self.categorical_colors = [
            "00D084",
            "00B4D8",
            "7B2CBF",
            "FF7E67",
            "F9A03F",
            "FF007F",
            "8A2BE2",
        ]

        self.vital_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        self.vital_font = Font(color="065F46", bold=True)
        self.trivial_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        self.trivial_font = Font(color="92400E", bold=False)


# ── Tokens de diseño para expedientes: REGLA ÚNICA para todos los casos ──────
# Base común (Calculo.xlsx, Informe.docx, Historico.xlsx, CalculoND):
#   navy + grises + semáforo (ok/warn/error/info) + escala 14/12/11/10/9/8.
# Familias: Calibri en Cálculo/Informe; Arial solo en Histórico clásico
# (situación legada ERP, deliberada).
# Variantes situacionales (mismo lenguaje, distinto uso):
#   meta cumplida/no (rebate), celdas editables ERP, caja de cálculo,
#   filas de notas en histórico, acento del reporte de evidencia.
EXP_NAVY = "0D2B4E"
EXP_BORDER = "BFBFBF"
EXP_BORDER_STRONG = "003366"
EXP_LABEL_BG = "F0F4FA"
EXP_TOTAL_BG = "FDE8E8"
EXP_WHITE = "FFFFFF"
EXP_INK = "1A1A1A"
EXP_INK_SOFT = "333333"
EXP_GRAY = "666666"
EXP_MUTED = "595959"
EXP_FAINT = "888888"
EXP_ERR_BG = "FFC7CE"
EXP_ERR_TX = "9C0006"
EXP_WARN_BG = "FFF2CC"
EXP_WARN_TX = "9C5700"
EXP_INFO_BG = "D9E1F2"
EXP_INFO_TX = "003366"
EXP_OK_BG = "C6EFCE"
EXP_OK_TX = "006100"
EXP_META_OK_BG = "E8F5E9"
EXP_META_FAIL_BG = "FFEBEE"
EXP_META_OK_TX = "2E7D32"
EXP_META_FAIL_TX = "C62828"
EXP_EDIT_BG = "E2EFDA"
EXP_HIGHLIGHT_BG = "FCE4D6"
EXP_NOTE_BG = "FFF8E1"
EXP_NOTE_BD = "F59E0B"
EXP_NOTE_TX = "92400E"
EXP_HIST_NOTA_BG = "E7E6E6"
EXP_ACCENT = "34D399"

EXP_T_BASE = 10


def exp_fill(hex6: str) -> PatternFill:
    """Relleno sólido openpyxl desde un token de la regla."""
    return PatternFill(start_color=hex6, end_color=hex6, fill_type="solid")


def exp_rgb(hex6: str):
    """RGBColor python-docx desde un token de la regla (import perezoso)."""
    from docx.shared import RGBColor

    h = hex6.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
