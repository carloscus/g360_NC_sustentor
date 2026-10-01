from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.comments import Comment
from src.render.g360_styles import G360Styles


_COLUMN_FORMATS = {
    "CODIGO_SKU": "@",
    "PRECIO_BASE": "#,##0.00000",
    "PRECIO_CORRECTO_SIN_IGV": "#,##0.00000",
    "PRECIO_UNITARIO_SIN_IGV": "#,##0.00000",
    "META_MONTO_SOLES": "#,##0.00",
    "PORCENTAJE_REBATE": "0.00%",
    "DESCUENTO_PORCENTAJE": "0.00%",
    "CANTIDAD_A_SUSTENTAR": "#,##0",
    "CANTIDAD_DEVUELTA": "#,##0",
    "FECHA_DEVOLUCION": "@",
    "UNIDADES_X_PRESENTACION": "#,##0",
    "CANTIDAD_STOCK": "#,##0",
    "DESC01": "0.00%",
    "DESC02": "0.00%",
    "DESC03": "0.00%",
    "DESC04": "0.00%",
    "DESC05": "0.00%",
    "DESC06": "0.00%",
    "DESC07": "0.00%",
    "DESC08": "0.00%",
    "REGLA_PROMOCIONAL": "@",
    "SUCURSAL": "@",
    "ARTICULO": "@",
}


class PlantillaGenerator:
    """Genera las plantillas descargables de insumos (una por archivo que se carga).

    Convención de nombres (única para lo generado y para lo que vive en
    ``assets/templates``), sin marca y en mayúsculas por palabra::

        Lista_de_<Precios|Descuentos|Devoluciones>[_y_Cantidades].xlsx

    - ``Precios`` / ``Descuentos`` / ``Devoluciones`` = la lista que el usuario
      reconoce.
    - ``_y_Cantidades`` = el archivo además trae la cantidad a reconocer.

    Las 5 listas cubren los insumos que se cargan como archivo:

    =================================  ====================  =========================
    Archivo                            Casos                Insumos
    =================================  ====================  =========================
    Lista_de_Precios                   DC                   lista_precios
    Lista_de_Precios_y_Cantidades      VRS                  lista_precios, cantidad
    Lista_de_Descuentos                DO                   porcentaje
    Lista_de_Descuentos_y_Cantidades   FPE (PROM/CMV       porcentaje, cantidad
                                       comparten columnas)
    Lista_de_Devoluciones              DF                   cantidad
    =================================  ====================  =========================

    Las columnas usan SIEMPRE el nombre que ve el usuario en el archivo
    (``CODIGO_SKU``, ``PRECIO_BASE``, ``CANTIDAD_A_SUSTENTAR``,
    ``DESCUENTO_PORCENTAJE``, ``DESC01..DESC08``); el ``header_map`` del
    proceso las lleva al nombre interno. El SKU es siempre ``CODIGO_SKU`` para
    que un mismo archivo sirva para todos los casos.
    """

    TEMPLATES = {
        "lista_precios": {
            "nombre": "Lista_de_Precios",
            "hoja": "LISTA_PRECIOS",
            "casos": ["DC"],
            "insumos": ["lista_precios"],
            "descripcion": "Precio de lista por SKU (DC). Las cantidades salen de las facturas.",
            "columnas": [
                ("CODIGO_SKU", "Código del Artículo / SKU"),
                ("PRECIO_BASE", "Precio base del producto sin IGV (ej: 5.14)"),
                ("DESC01", "Primer descuento en % (fracción: 0.25 = 25%)"),
                ("DESC02", "Segundo descuento en % (fracción)"),
                ("DESC03", "Tercer descuento en % (fracción)"),
                ("DESC04", "Cuarto descuento en % (fracción)"),
                ("DESC05", "Quinto descuento en % (fracción)"),
                ("DESC06", "Sexto descuento en % (fracción)"),
                ("DESC07", "Séptimo descuento en % (fracción)"),
                ("DESC08", "Octavo descuento en % (fracción)"),
            ],
            "ejemplos": [
                ["03108", 5.14, 0.25, 0.04, 0.04, 0.07, 0.02, 0, 0, 0],
                ["A002", 3.50, 0.10, 0.05, 0, 0, 0, 0, 0, 0],
            ],
        },
        "lista_precios_cantidad": {
            "nombre": "Lista_de_Precios_y_Cantidades",
            "hoja": "LISTA_PRECIOS_CANT",
            "casos": ["VRS"],
            "insumos": ["lista_precios", "cantidad"],
            "descripcion": "Precio de lista y cantidad a reconocer por SKU (VRS).",
            "columnas": [
                ("CODIGO_SKU", "Código del Artículo / SKU"),
                ("CANTIDAD_A_SUSTENTAR", "Cantidad de unidades a reconocer por SKU"),
                ("PRECIO_BASE", "Precio de lista vigente sin IGV antes de descuentos"),
                ("DESC01", "Descuento 1 (fracción, ej: 0.05 = 5%) — opcional"),
                ("DESC02", "Descuento 2 (fracción) — opcional"),
            ],
            "ejemplos": [
                ["72015", 12954, 4.80, 0.05, 0],
                ["72015", 500, 4.80, 0, 0.05],
                ["A002", 500, 12.00, None, None],
            ],
        },
        "lista_descuentos": {
            "nombre": "Lista_de_Descuentos",
            "hoja": "LISTA_DESCUENTOS",
            "casos": ["DO"],
            "insumos": ["porcentaje"],
            "descripcion": "Descuento por SKU (DO). Se aplica sobre el precio atendido.",
            "columnas": [
                ("CODIGO_SKU", "Código del Artículo / SKU a descontar"),
                (
                    "DESCUENTO_PORCENTAJE",
                    "Porcentaje de descuento (0.05 = 5%; también acepta 5 = 5%)",
                ),
            ],
            "ejemplos": [
                ["A001", 0.05],
                ["A002", 0.03],
                ["A003", 0.025],
            ],
        },
        "lista_descuentos_cantidad": {
            "nombre": "Lista_de_Descuentos_y_Cantidades",
            "hoja": "LISTA_DESCUENTOS_CANT",
            "casos": ["FPE", "PROM", "CMV"],
            "insumos": ["porcentaje", "cantidad"],
            "descripcion": "Cantidad a sustentar y descuento por SKU (FPE; columnas compatibles con PROM y CMV).",
            "columnas": [
                ("CODIGO_SKU", "Código del Artículo / SKU a sustentar"),
                ("CANTIDAD_A_SUSTENTAR", "Cantidad total de unidades a sustentar"),
                ("DESCUENTO_PORCENTAJE", "Porcentaje de descuento (0.05 = 5%)"),
            ],
            "ejemplos": [
                ["A001", 100, 0.05],
                ["A002", 50, 0.10],
            ],
        },
        "lista_devoluciones": {
            "nombre": "Lista_de_Devoluciones",
            "hoja": "LISTA_DEVOLUCIONES",
            "casos": ["DF"],
            "insumos": ["cantidad"],
            "descripcion": "Unidades devueltas por SKU (DF). Se asignan LIFO contra las facturas.",
            "columnas": [
                ("CODIGO_SKU", "Código del Artículo / SKU devuelto"),
                ("CANTIDAD_DEVUELTA", "Unidades devueltas del SKU"),
                (
                    "FECHA_DEVOLUCION",
                    "Fecha de la devolución (opcional; si se omite se toman todas las facturas)",
                ),
            ],
            "ejemplos": [
                ["03108", 20, "2026-02-01"],
                ["72015", 5, None],
            ],
        },
    }

    # Claves de la versión anterior (por proceso, no por insumo). Se aceptan al
    # generar para no romper scripts externos; la clave vigente es la de arriba.
    CLAVES_LEGACY = {
        "feria_preventa": "lista_descuentos_cantidad",
        "descuento_sku": "lista_descuentos",
        "bonificacion_rebate": "lista_descuentos_cantidad",
        "stock_cliente": "lista_precios_cantidad",
    }

    @classmethod
    def resolver_clave(cls, tipo: str) -> str:
        """Clave vigente para un tipo dado (acepta las claves legacy)."""
        return cls.CLAVES_LEGACY.get(tipo, tipo)

    def __init__(self):
        self.styles = G360Styles()

    def generar(self, tipo: str, ruta_salida: str) -> Path:
        clave = self.resolver_clave(tipo)
        template = self.TEMPLATES.get(clave)
        if not template:
            raise ValueError(
                f"Plantilla desconocida: {tipo}. Disponibles: {', '.join(self.TEMPLATES)}"
            )

        out_path = Path(ruta_salida)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        wb.properties.creator = "ccusi"
        wb.properties.description = f"{template['nombre']} — {template['descripcion']}"
        ws = wb.active
        ws.title = template.get("hoja") or template["nombre"][:31]

        columnas = template["columnas"]
        ejemplos = template["ejemplos"]

        # Escribir cabeceras con comentario y formato
        for col, (nombre, descripcion) in enumerate(columnas, 1):
            celda = ws.cell(row=1, column=col, value=nombre)
            celda.fill = self.styles.header_fill
            celda.font = self.styles.header_font
            celda.alignment = self.styles.center_align
            celda.border = self.styles.border
            celda.comment = Comment(descripcion, "ccusi")
            fmt = _COLUMN_FORMATS.get(nombre)
            if fmt:
                celda.number_format = fmt

        # Escribir ejemplos con formato numérico
        for fila, datos in enumerate(ejemplos, 2):
            for col, valor in enumerate(datos, 1):
                celda = ws.cell(row=fila, column=col, value=valor if valor != "" else None)
                celda.fill = self.styles.zebra_fill
                celda.border = self.styles.border
                nombre = columnas[col - 1][0]
                fmt = _COLUMN_FORMATS.get(nombre)
                if fmt:
                    celda.number_format = fmt

        # Instrucciones en una hoja aparte (LEEME).
        # Si se escriben en la MISMA hoja, `read_erp_file` las devuelve como
        # filas de datos: FPE las reporta como "SKU no encontrado" (AL06) y el
        # resto de los casos las arrastra al preprocessado.
        ws_leeme = wb.create_sheet("LEEME")
        ws_leeme.column_dimensions["A"].width = 78
        celda_nota = ws_leeme.cell(row=1, column=1, value="INSTRUCCIONES:")
        celda_nota.font = Font(bold=True, size=11)
        celda_nota.fill = self.styles.note_fill
        celda_desc = ws_leeme.cell(row=2, column=1, value=template.get("descripcion", ""))
        celda_desc.font = Font(size=10, italic=True, color="444444")
        # A qué casos e insumos sirve: la convención de nombres tiene que ser
        # legible sin abrir la app.
        celda_uso = ws_leeme.cell(
            row=3,
            column=1,
            value=(
                f"Casos: {', '.join(template.get('casos', []))} · "
                f"Insumos: {', '.join(template.get('insumos', []))}"
            ),
        )
        celda_uso.font = Font(size=10, color="444444")

        instrucciones = [
            "1. Eliminar las filas de ejemplo antes de cargar tus datos",
            "2. No modificar el nombre ni orden de las columnas",
            "3. Los descuentos se ingresan en formato decimal (0.10 = 10%)",
            "4. No dejar filas vacías entre registros",
            "5. Guardar el archivo antes de importar al sistema",
        ]
        for i, texto in enumerate(instrucciones, 5):
            ws_leeme.cell(row=i, column=1, value=texto).font = Font(size=10, color="444444")

        # Ajustar anchos
        anchos = {1: 18, 2: 18, 3: 18, 4: 18, 5: 14, 6: 14, 7: 14, 8: 18, 9: 16, 10: 16}
        for col, ancho in anchos.items():
            if col <= len(columnas):
                ws.column_dimensions[chr(64 + col)].width = ancho

        ws.freeze_panes = "A2"
        if columnas:
            ws.auto_filter.ref = f"A1:{chr(64 + len(columnas))}1"
        try:
            wb.save(str(out_path))
        except PermissionError:
            raise PermissionError(
                f"No se pudo guardar la plantilla. ¿Está abierto en otro programa?\n"
                f"Cierre el archivo e intente nuevamente.\nRuta: {out_path}"
            )
        return out_path
