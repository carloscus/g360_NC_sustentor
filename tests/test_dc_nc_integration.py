"""NC/NDB reconciliation semantics in the Difference de Precio engine."""

import pandas as pd
from openpyxl import load_workbook

from src.core.nc_reconciliation import reconciliar_notas
from src.domain import ExpedienteComercial, PipelineContext, RecognitionResult
from src.render.excel_renderer import ExcelRenderer
from src.strategies.price_difference import PriceDifferenceStrategy
from src.pipeline import Pipeline
from src.ui.config_builder import build_config, build_datos_exp


def _history(note_rows=(), extra_invoice=False):
    rows = [
        {
            "CODIGO": "00123",
            "ARTICULO": "Producto X",
            "CANTIDAD": 100.0,
            "SOLES": 1000.0,
            "FECHA": pd.Timestamp("2026-01-15"),
            "TIPO_DOC": "F01",
            "SERIE": "001",
            "NUMERO": "10",
            "DOC_ID": "F001-10",
            "TIPO_CLASE": "factura",
            "FACTURA_REF": "",
            "CANTIDAD_FAE": 0.0,
        }
    ]
    if extra_invoice:
        rows.append(
            {
                **rows[0],
                "SERIE": "002",
                "NUMERO": "20",
                "DOC_ID": "F002-20",
                "FECHA": pd.Timestamp("2026-01-16"),
                "SOLES": 1200.0,
            }
        )
    for i, n in enumerate(note_rows, 1):
        rows.append(
            {
                "CODIGO": n.get("sku", "00123"),
                "ARTICULO": "Producto X",
                "CANTIDAD": n.get("cantidad", 0.0),
                "CANTIDAD_FAE": n.get("fae", 0.0),
                "SOLES": n.get("soles", -100.0),
                "FECHA": pd.Timestamp("2026-01-20"),
                "TIPO_DOC": n.get("tipo", "NCR"),
                "SERIE": "N900",
                "NUMERO": str(i),
                "DOC_ID": f"NCRN900-{i}",
                "TIPO_CLASE": n.get("clase", "descuento"),
                "FACTURA_REF": n.get("ref", "F001-10"),
                "REFERENCIA": "F01/001-10",
            }
        )
    return pd.DataFrame(rows)


def _run(hist, modalidad="individual", use_nc=True, use_ndb=True):
    options = {
        "facturas": {"mostrar": True, "usar": True},
        "nc": {"mostrar": True, "usar": use_nc},
        "ndb": {"mostrar": True, "usar": use_ndb},
    }
    rec = reconciliar_notas(hist, documentos=options, modalidad=modalidad)
    facts = hist[hist.TIPO_CLASE == "factura"].copy()
    facts["DOC_ID"] = facts.apply(lambda r: f"{r.TIPO_DOC[0]}{r.SERIE}-{r.NUMERO}", axis=1)
    lista = pd.DataFrame({"SKU": ["00123"], "PRECIO_BASE": [8.0]})
    result, row_alerts, full = PriceDifferenceStrategy()._procesar_comparar(
        facts,
        lista,
        [],
        config={"modalidad": modalidad, "sort_mode": "fecha_asc", "reconciliacion_nc": rec},
    )
    return result, row_alerts, full


def test_exact_fae_changes_effective_price_but_not_invoice_baseline():
    hist = _history([{"fae": 100, "soles": -100}])
    result, _alerts, full = _run(hist)
    row = result.iloc[0]
    assert row["PRECIO_HIST"] == 9.0
    assert row["PRECIO_HIST_FACTURA"] == 10.0
    assert row["SOLES"] == 1000.0
    assert row["SOLES_CALCULO"] == 900.0
    assert row["MONTO_NC"] == 100.0
    assert "APLICADA" in row["ALERTA"]
    assert "N900-1" in row["AUDITORIA_NC"]
    assert full.iloc[0]["TOTAL_FACTURA_EXACTO"] == 1000.0


def test_exact_return_reduces_eligible_qty_not_unit_price():
    hist = _history([{"clase": "devolucion", "cantidad": -100, "soles": -1000}])
    result, _alerts, _full = _run(hist)
    row = result.iloc[0]
    assert row["CANTIDAD_FACTURADA"] == 100
    assert row["CANTIDAD"] == 0
    assert row["PRECIO_HIST"] == 10
    assert row["MONTO_NC"] == 0
    assert "cantidad reducida" in row["ALERTA"]


def test_partial_and_excess_notes_do_not_apply_but_are_alerted():
    partial, _, _ = _run(_history([{"fae": 40, "soles": -40}]))
    assert partial.iloc[0]["SOLES_CALCULO"] == 1000
    assert partial.iloc[0]["CANTIDAD"] == 100
    assert partial.iloc[0]["MONTO_NC"] == 200
    assert "REVISIÓN MANUAL" in partial.iloc[0]["ALERTA"]

    excess, _, _ = _run(_history([{"fae": 120, "soles": -120}]))
    assert excess.iloc[0]["SOLES_CALCULO"] == 1000
    assert excess.iloc[0]["CANTIDAD"] == 100
    assert excess.iloc[0]["MONTO_NC"] == 200
    assert "AL13" in excess.iloc[0]["ALERTA"]
    assert "no aplicado" in excess.iloc[0]["ALERTA"]


def test_consolidated_notes_are_detailed_but_not_applied_and_keep_all_invoices():
    hist = _history(
        [
            {"fae": 50, "soles": -50, "ref": "F001-10"},
            {"fae": 50, "soles": -50, "ref": "F002-20"},
        ],
        extra_invoice=True,
    )
    result, alerts, _full = _run(hist, modalidad="consolidado")
    row = result.iloc[0]
    assert row["CANTIDAD"] == 200
    assert row["SOLES_CALCULO"] == 2200
    assert row["MONTO_NC"] == 600
    assert row["FACTURAS"] == "F001-10, F002-20"
    assert "no aplicado" in row["ALERTA"]
    assert "N900-1" in row["AUDITORIA_NC"]
    assert "N900-2" in row["AUDITORIA_NC"]
    assert len([a for a in alerts if a.codigo == "AL01"]) == 1


def test_exact_ndb_fae_applies_signed_charge_to_historical_value():
    hist = _history([{"tipo": "NDB", "clase": "cargo", "fae": 100, "soles": 100}])
    result, _alerts, _full = _run(hist)
    row = result.iloc[0]
    assert row["SOLES_CALCULO"] == 1100
    assert row["PRECIO_HIST"] == 11
    assert row["MONTO_NC"] == 300


def test_excel_calculo_uses_effective_price_and_keeps_note_audit(tmp_path):
    hist = _history([{"fae": 100, "soles": -100}])
    result, _alerts, full = _run(hist)
    wrapped = RecognitionResult(
        dataframe=result,
        dataframe_excel=full,
        resumen={
            "total_nc": float(result["MONTO_NC"].sum()),
            "skus_afectados": 1,
            "doc_ref": "F001-10",
        },
    )
    out = tmp_path / "DC_fae_exacta.xlsx"
    ExcelRenderer().generar(
        wrapped,
        str(out),
        tipo="diferencia_precio",
        cliente="CLIENTE X",
        doc_ref="F001-10",
        df_historial=hist,
        modalidad="individual",
    )

    wb = load_workbook(out, read_only=True, data_only=False)
    ws = wb[wb.sheetnames[0]]
    header = next(
        r for r in range(1, ws.max_row + 1) if "PRECIO HIST. EFECTIVO" in [c.value for c in ws[r]]
    )
    headers = [c.value for c in ws[header]]
    row = [c.value for c in ws[header + 1]]
    assert row[headers.index("PRECIO HIST. EFECTIVO")] == 9.0
    assert "N900-1" in row[headers.index("AUDITORÍA")]
    assert row[headers.index("DIF. UNITARIA")].startswith("=MAX(0,ROUND(")
    assert row[headers.index("MONTO")].startswith("=ROUND(")
    wb.close()


def test_pipeline_uses_live_history_checks_for_exact_note():
    hist = _history([{"fae": 100, "soles": -100}])
    policy = {
        "facturas": [True, True],
        "nc": [True, True],
        "ndb": [True, False],
    }
    ui = {
        "modalidad": "individual",
        "sort_mode_dc": "fecha_asc",
        "historico_config": policy,
        "cliente_pb": "CLIENTE X",
        "df_historial_full": hist,
    }
    config = build_config("diferencia_precio", ui)
    data = build_datos_exp("diferencia_precio", hist, config, ui)
    exp = ExpedienteComercial(
        nombre="Diferencia de precio",
        familia="precio",
        estrategia="PriceDifference",
        datos=data,
        contexto=PipelineContext(config=config),
    )
    exp.condiciones = [pd.DataFrame({"SKU": ["00123"], "PRECIO_BASE": [8.0]})]
    exp = Pipeline().ejecutar(exp)
    assert exp.resultado is not None
    row = exp.resultado.dataframe.iloc[0]
    assert row["SOLES_CALCULO"] == 900.0
    assert row["PRECIO_HIST"] == 9.0
    assert row["MONTO_NC"] == 100.0
    assert "APLICADA" in row["ALERTA"]
class TestReferenciaSinLetraDeTipo:
    """El folio con SKU y numero tiene que matchear aunque falte la letra.

    El ERP exporta la referencia a veces como "F204-67630" y a veces como
    "204-67630". Antes, sin letra no matcheaba con el DOC_ID (que la lleva)
    y la NC caia a revision manual aunque fuera exacta. Ahora hay fallback
    insensible a la letra, solo si es inequivoco (un unico candidato).
    """

    def _hist(self, ref_col, ref_val):
        base = {
            "CODIGO": "09008",
            "CANTIDAD": 720.0,
            "SOLES": -1800.0,
            "TIPO_DOC": "NCR",
            "SERIE": "215",
            "NUMERO": "31630",
            "DOC_ID": "NCR215-31630",
            "TIPO_CLASE": "nota_credito",
            "COD_CLIENTE": "C1",
            "FACTURA_REF": "",
            "FACTURA_REF_SERIE": "",
            "FACTURA_REF_NRO": "",
            "REFERENCIA": "",
        }
        base[ref_col] = ref_val
        fac = {
            "CODIGO": "09008",
            "CANTIDAD": 720.0,
            "SOLES": 1800.0,
            "TIPO_DOC": "F01",
            "SERIE": "204",
            "NUMERO": "67630",
            "DOC_ID": "F204-67630",
            "TIPO_CLASE": "factura",
            "FACTURA_REF": "",
            "FACTURA_REF_SERIE": "",
            "FACTURA_REF_NRO": "",
            "REFERENCIA": "",
            "COD_CLIENTE": "C1",
        }
        import pandas as pd

        return pd.DataFrame([fac, base])

    def _options(self):
        return {
            "facturas": {"mostrar": True, "usar": True},
            "nc": {"mostrar": True, "usar": True},
            "ndb": {"mostrar": True, "usar": False},
        }

    def _rec(self, ref_col, ref_val):
        from src.core.nc_reconciliation import reconciliar_notas

        rec = reconciliar_notas(
            self._hist(ref_col, ref_val),
            documentos=self._options(),
            modalidad="individual",
        )
        assert len(rec) == 1
        return list(rec.values())[0]

    def test_con_letra_aplica(self):
        v = self._rec("FACTURA_REF", "F204-67630")
        assert v["qty_return"] == 720.0
        assert v["applied"] is True

    def test_sin_letra_aplica_igual(self):
        v = self._rec("FACTURA_REF", "204-67630")
        assert v["qty_return"] == 720.0, "el fallback tiene que rescatarla"
        assert v["applied"] is True

    def test_referencia_libre_aplica(self):
        v = self._rec("REFERENCIA", "Devolucion factura 204-67630 por reclamo")
        assert v["qty_return"] == 720.0
        assert v["applied"] is True

    def test_avisa_que_fue_aproximada(self):
        v = self._rec("FACTURA_REF", "204-67630")
        textos = " ".join(v.get("alerts", []))
        assert "ref. aproximada" in textos
        v2 = self._rec("FACTURA_REF", "F204-67630")
        textos2 = " ".join(v2.get("alerts", []))
        assert "ref. aproximada" not in textos2, "con letra no hay nada que avisar"

    def test_ambiguo_no_adivina(self):
        """Si existen F204-67630 y B204-67630, '204-67630' no elige solo."""
        import pandas as pd

        df = self._hist("FACTURA_REF", "204-67630")
        extra = pd.DataFrame(
            [
                {
                    "CODIGO": "09008",
                    "CANTIDAD": 720.0,
                    "SOLES": 1800.0,
                    "TIPO_DOC": "F02",
                    "SERIE": "B04",
                    "NUMERO": "67630",
                    "DOC_ID": "B204-67630",
                    "TIPO_CLASE": "factura",
                    "FACTURA_REF": "",
                    "FACTURA_REF_SERIE": "",
                    "FACTURA_REF_NRO": "",
                    "REFERENCIA": "",
                    "COD_CLIENTE": "C1",
                }
            ]
        )
        df = pd.concat([df, extra], ignore_index=True)
        from src.core.nc_reconciliation import reconciliar_notas

        rec = reconciliar_notas(df, documentos=self._options(), modalidad="individual")
        # Ninguna de las dos facturas debe quedar aplicada por adivinanza.
        aplicadas = [v for v in rec.values() if v.get("applied")]
        assert aplicadas == [], "ambiguo: mejor manual que mal aplicado"

class TestNotaHuerfana:
    """NC que referencia factura conocida pero SKU ajeno: avisa, no calcula.

    Antes se perdia en silencio: ni entraba al calculo ni avisaba. Ahora
    deja una alerta AJENA con cantidad y monto en la entrada de la factura,
    sin entrar a preview ni a calculos. El monto importa para no exceder el
    total de la factura con NCs futuras.
    """

    def _hist_huerfana(self, sku="99999", cant=10.0, soles=-25.0):
        import pandas as pd

        return pd.DataFrame(
            [
                {
                    "CODIGO": "09008",
                    "CANTIDAD": 720.0,
                    "SOLES": 1800.0,
                    "TIPO_DOC": "F01",
                    "SERIE": "204",
                    "NUMERO": "67630",
                    "DOC_ID": "F204-67630",
                    "TIPO_CLASE": "factura",
                    "FACTURA_REF": "",
                    "COD_CLIENTE": "C1",
                    "CANTIDAD_FAE": 0.0,
                },
                {
                    "CODIGO": sku,
                    "CANTIDAD": cant,
                    "SOLES": soles,
                    "TIPO_DOC": "NCR",
                    "SERIE": "215",
                    "NUMERO": "31630",
                    "DOC_ID": "NCR215-31630",
                    "TIPO_CLASE": "nota_credito",
                    "FACTURA_REF": "F204-67630",
                    "COD_CLIENTE": "C1",
                    "CANTIDAD_FAE": 0.0,
                },
            ]
        )

    def _options(self):
        return {
            "facturas": {"mostrar": True, "usar": True},
            "nc": {"mostrar": True, "usar": True},
            "ndb": {"mostrar": True, "usar": False},
        }

    def test_avisa_con_monto_y_no_aplica(self):
        from src.core.nc_reconciliation import reconciliar_notas

        rec = reconciliar_notas(
            self._hist_huerfana(), documentos=self._options(), modalidad="individual"
        )
        assert len(rec) == 1
        v = list(rec.values())[0]
        assert v.get("applied") is False
        textos = " ".join(v.get("alerts", []))
        assert "AJENA" in textos
        assert "99999" in textos
        assert "25.00" in textos, "el monto tiene que estar para no exceder el total"

    def test_no_mete_el_sku_ajeno_al_calculo(self):
        """El 99999 no puede aparecer como linea calculable."""
        from src.core.nc_reconciliation import reconciliar_notas

        rec = reconciliar_notas(
            self._hist_huerfana(), documentos=self._options(), modalidad="individual"
        )
        assert all("99999" not in str(k) for k in rec.keys())

    def test_dos_parciales_no_aplican(self):
        """Dos devoluciones parciales no se combinan para aplicar."""
        import pandas as pd

        from src.core.nc_reconciliation import reconciliar_notas

        base = self._hist_huerfana(sku="09008", cant=100.0, soles=-250.0)
        extra = pd.DataFrame(
            [
                {
                    "CODIGO": "09008",
                    "CANTIDAD": 100.0,
                    "SOLES": -250.0,
                    "TIPO_DOC": "NCR",
                    "SERIE": "215",
                    "NUMERO": "2",
                    "DOC_ID": "NCR215-2",
                    "TIPO_CLASE": "nota_credito",
                    "FACTURA_REF": "F204-67630",
                    "COD_CLIENTE": "C1",
                    "CANTIDAD_FAE": 0.0,
                }
            ]
        )
        hist = pd.concat([base, extra], ignore_index=True)
        rec = reconciliar_notas(hist, documentos=self._options(), modalidad="individual")
        assert all(v.get("applied") is not True for v in rec.values())

