"""El Histórico del expediente: un solo cliente, su período, y resaltado de lo usado.

El Histórico se arma desde el fragmento del cliente (período completo, líneas
del allowlist). Lo que marca el cambio de cada expediente es el conjunto de
pares (factura, SKU) que ese expediente realmente usó: esas filas quedan con
el relleno de reclamo (FFF2CC) y las notas con su propio color.

Este test lo verifica sobre los XLSX generados, no sobre funciones.
"""

import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).parent.parent))

from src.core import ventas_db
from src.core.ventas_db_client import VentasDbClient
from src.core.xls_processor import derivar_campos
from src.ui.expediente_service import generar_expediente
from src.ui.reconocimiento_view import ReconocimientoView

from tests.conftest import *  # noqa: F401,F403  (fixtures: tmp_db)
from tests.test_render_parentdata import _AppStub  # noqa: E402

ALFA = ("00056101", "20601024714", "CLIENTE ALFA S.A.")
BETA = ("00005610", "20601024715", "CLIENTE BETA E.I.R.L.")

RECLAMO_FILL = "FFF2CC"
COLS = {  # posiciones del encabezado del Historico
    "FECHA": 1,
    "TIPO": 2,
    "SERIE": 3,
    "NUMERO": 4,
    "SKU": 5,
    "RUC": 10,
}


def _fila(cid, ruc, nom, sku, tpo, serie, nro, soles, cantidad, fecha, **extra):
    d = dict(
        id_articulo=sku,
        original_sku=sku,
        nom_articulo=f"ART {sku}",
        id_linea="0101",
        nom_linea="GASEOSAS",
        id_grupo="01",
        nom_grupo="G",
        id_tipo="01",
        nom_tipo="T",
        id_familia="01",
        nom_familia="F",
        id_cliente=cid,
        doc_cliente=ruc,
        nom_cliente=nom,
        tpo_doc=tpo,
        serie_doc=serie,
        nro_doc=nro,
        referencia="",
        moneda="Soles",
        cantidad=cantidad,
        cantidad_fae=0.0,
        soles=soles,
        dolares=0.0,
        precio_unitario=soles / cantidad if cantidad else 0.0,
        anho=2026,
        mes=9,
        fecha_orig=fecha,
        fecha_ref=None,
        fecha_venc=None,
        cod_sucursal="01",
        nom_sucursal="LIMA",
        departamento="LIMA",
        provincia="LIMA",
        distrito="SAN ISIDRO",
        id_vendedor="01188",
        nom_vendedor="VEND",
        id_pedido="P1",
        file_source="t",
        mes_ref="2026-09",
        tipo_operacion="",
        factura_ref_serie="",
        factura_ref_nro="",
        folio_unico="",
    )
    d.update(extra)
    return d


def _poblar():
    """ALFA: 3 facturas (2 SKU) + 1 NC. BETA: 2 facturas + 1 NC."""
    filas = []
    n = [67370]
    plan = [
        (
            ALFA,
            [
                ("014850", 500.0, 100.0, "2026-09-05"),
                ("014850", 380.0, 80.0, "2026-09-12"),
                ("031080", 240.0, 100.0, "2026-09-18"),
            ],
        ),
        (
            BETA,
            [
                ("014851", 980.0, 100.0, "2026-09-07"),
                ("014851", 490.0, 50.0, "2026-09-22"),
            ],
        ),
    ]
    for (cid, ruc, nom), items in plan:
        primera = None
        for sku, soles, cant, fecha in items:
            n[0] += 1
            filas.append(_fila(cid, ruc, nom, sku, "F01", "204", str(n[0]), soles, cant, fecha))
            if primera is None:
                primera = n[0]
        filas.append(
            _fila(
                cid,
                ruc,
                nom,
                items[0][0],
                "NCR",
                "N204",
                f"9{primera}",
                -round(items[0][1] * 0.2, 2),
                -round(items[0][2] * 0.2, 2),
                "2026-09-25",
                referencia=f"F01/204-{primera}",
                factura_ref_serie="204",
                factura_ref_nro=str(primera),
            )
        )
    rows = []
    for d in filas:
        d = dict(d)
        derivar_campos(d)
        rows.append(d)
    conn = ventas_db.get_conn()
    ventas_db.insert_ventas(conn, rows)
    ventas_db.populate_nc_asociadas(conn)


def _historial(clientes):
    cli = VentasDbClient()
    partes = [
        cli.fetch_historial(id_cliente=cid, fecha_desde="2026-09-01", fecha_hasta="2026-09-30")
        for cid, _r, _n in clientes
    ]
    return pd.concat(partes, ignore_index=True)


def _skus(clientes):
    return {ALFA[0]: ["014850", "031080"], BETA[0]: ["014851"]}


def _vista(tmp_path, clientes, modalidad, destino):
    app = _AppStub()
    errores = []
    app.show_snackbar = lambda msg, *a, **k: errores.append(msg)
    view = ReconocimientoView(app)
    view.build()
    view.tipo_actual = "DC"
    view.tipo_dropdown.value = "DC"
    view.modalidad_radio.value = modalidad
    view.df_historial = _historial(clientes)
    lista = tmp_path / f"lista_{destino}.csv"
    lista.write_text(
        "CODIGO_SKU,PRECIO_BASE,DESC01\n014850,4.20,0.05\n014851,8.90,0.10\n",
        encoding="utf-8",
    )
    view.lista_path = str(lista)
    view._ejecutar_tarea()
    assert errores == [], errores
    return view


def _generar(view, tmp_path, modalidad, sufijo):
    from src.ui.catalog import HistorialConfig

    dest = tmp_path / f"exp_{sufijo}"
    dest.mkdir()
    return generar_expediente(
        resultado=view.resultado,
        tipo_actual="DC",
        modalidad=modalidad,
        historico_config=HistorialConfig(facturas=(True, True), nc=(True, True), ndb=(True, False)),
        df_historial=view.df_historial,
        cliente_value="",
        vendedor_value="",
        vendedor_display="VEND",
        antecedentes="sim historico",
        observaciones="",
        desktop_path=dest,
    )


def _leer_historico(exp_dir):
    """Del XLSX del Historico: rucs, periodo y filas resaltadas (factura, sku)."""
    from openpyxl import load_workbook

    ruta = next(exp_dir.glob("*Historico.xlsx"))
    wb = load_workbook(str(ruta))
    try:
        info = {"rucs": set(), "fechas": [], "resaltadas": set(), "notas": set(), "filas": 0}
        for hoja in wb.sheetnames:
            ws = wb[hoja]
            hdr = None
            for r in range(1, ws.max_row + 1):
                vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                limpios = [str(v).strip() for v in vals if v is not None]
                if hdr is None:
                    if "SKU" in limpios and "SERIE" in limpios:
                        hdr = r
                    continue
                tipo = str(ws.cell(row=r, column=COLS["TIPO"]).value or "").strip()
                serie = str(ws.cell(row=r, column=COLS["SERIE"]).value or "").strip()
                numero = str(ws.cell(row=r, column=COLS["NUMERO"]).value or "").strip()
                sku = str(ws.cell(row=r, column=COLS["SKU"]).value or "").strip()
                fecha = ws.cell(row=r, column=COLS["FECHA"]).value
                ruc = ws.cell(row=r, column=COLS["RUC"]).value
                if ruc:
                    info["rucs"].add(str(ruc).strip())
                if fecha is not None:
                    info["fechas"].append(fecha)
                if not sku:
                    continue
                info["filas"] += 1
                fill = ws.cell(row=r, column=1).fill
                rgb = str(getattr(fill.start_color, "rgb", None) or "")
                fac = f"{tipo[:1]}{serie}-{numero}"
                if rgb.endswith(RECLAMO_FILL):
                    info["resaltadas"].add((fac, sku))
                elif tipo.upper().startswith("N"):
                    info["notas"].add((fac, sku))
        return info
    finally:
        wb.close()


def _usadas(view, fac_nro):
    df = view.resultado.resultado.dataframe
    sub = df[df["FACTURA"].astype(str).str.endswith(str(fac_nro))]
    return {(str(r["FACTURA"]).strip(), str(r["SKU"]).strip()) for _, r in sub.iterrows()}


def test_historico_individual_1_cliente(tmp_db, tmp_path):
    _poblar()
    view = _vista(tmp_path, [ALFA], "individual", "ind1")
    dirs = _generar(view, tmp_path, "individual", "ind1")
    print("\n### individual · 1 cliente ###")
    assert len(dirs) == 2
    for d in dirs:
        info = _leer_historico(d)
        usadas = _usadas(view, d.name.split("-")[4])
        print(
            "  %-42s ruc=%s filas=%d resaltadas=%s usadas=%s"
            % (
                d.name,
                sorted(info["rucs"]),
                info["filas"],
                sorted(info["resaltadas"]),
                sorted(usadas),
            )
        )
        assert info["rucs"] == {ALFA[1]}, info["rucs"]
        assert info["filas"] == 4, info["filas"]
        assert info["resaltadas"] == usadas, (info["resaltadas"], usadas)
        assert info["notas"], "la NC debia quedar con su relleno"
        assert not (info["notas"] & info["resaltadas"])


def test_historico_individual_2_clientes(tmp_db, tmp_path):
    _poblar()
    view = _vista(tmp_path, [ALFA, BETA], "individual", "ind2")
    dirs = _generar(view, tmp_path, "individual", "ind2")
    print("\n### individual · 2 clientes ###")
    assert len(dirs) == 4
    rucs = {ALFA[1]: 0, BETA[1]: 0}
    for d in dirs:
        info = _leer_historico(d)
        usadas = _usadas(view, d.name.split("-")[4])
        ruc = sorted(info["rucs"])[0]
        rucs[ruc] += 1
        print(
            "  %-42s ruc=%s filas=%d resaltadas=%s usadas=%s"
            % (d.name, ruc, info["filas"], sorted(info["resaltadas"]), sorted(usadas))
        )
        assert len(info["rucs"]) == 1, f"historico mezclando clientes: {info['rucs']}"
        assert info["resaltadas"] == usadas, (info["resaltadas"], usadas)
    assert rucs == {ALFA[1]: 2, BETA[1]: 2}, rucs


def test_historico_consolidado_2_clientes(tmp_db, tmp_path):
    _poblar()
    view = _vista(tmp_path, [ALFA, BETA], "consolidado", "cons")
    dirs = _generar(view, tmp_path, "consolidado", "cons")
    print("\n### consolidado · 2 clientes ###")
    assert len(dirs) == 2
    for d in dirs:
        info = _leer_historico(d)
        print(
            "  %-42s ruc=%s filas=%d resaltadas=%s"
            % (d.name, sorted(info["rucs"]), info["filas"], sorted(info["resaltadas"]))
        )
        assert len(info["rucs"]) == 1, f"historico mezclando clientes: {info['rucs']}"
        assert len(info["resaltadas"]) >= 1
        assert not (info["notas"] & info["resaltadas"])
