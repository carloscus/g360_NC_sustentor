"""End-to-end: la regla de NC sobrevive hasta el expediente en disco.

- Cálculo.xlsx: ALERTA (AL12) + AUDITORÍA con el folio de la nota.
- Histórico.xlsx: muestra factura y NC según "Mostrar" (independiente de "Usar").
- Informe.docx: se genera junto a los otros dos.
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.core.nc_auditor import CreditNoteAuditor
from src.domain import ExpedienteComercial, PipelineContext
from src.pipeline import Pipeline
from src.ui.catalog import HistorialConfig
from src.ui.config_builder import build_config, build_datos_exp
from src.ui.expediente_service import generar_expediente


def _history():
    return pd.DataFrame(
        [
            {
                "CODIGO": "00123",
                "ARTICULO": "Producto X",
                "CANTIDAD": 100.0,
                "SOLES": 1000.0,
                "FECHA": pd.Timestamp("2026-01-15"),
                "TIPO_DOC": "F01",
                "SERIE": "001",
                "NUMERO": "10",
                "DOC_ID": "F001-10",
                "TIPO_CLASE": "factura",
                "FACTURA_REF": "",
                "CANTIDAD_FAE": 0.0,
                "CLIENTE": "CLIENTE X",
                "COD_CLIENTE": "00056101",
                "DOC_CLIENTE": "206039928",
            },
            {
                "CODIGO": "00123",
                "ARTICULO": "Producto X",
                "CANTIDAD": 0.0,
                "CANTIDAD_FAE": 100.0,
                "SOLES": -100.0,
                "FECHA": pd.Timestamp("2026-01-20"),
                "TIPO_DOC": "NCR",
                "SERIE": "N900",
                "NUMERO": "1",
                "DOC_ID": "NCRN900-1",
                "TIPO_CLASE": "descuento",
                "FACTURA_REF": "F001-10",
                "REFERENCIA": "F01/001-10",
                "CLIENTE": "CLIENTE X",
                "COD_CLIENTE": "00056101",
                "DOC_CLIENTE": "206039928",
            },
        ]
    )


def _run_pipeline(hist, policy):
    ui = {
        "modalidad": "individual",
        "sort_mode_dc": "fecha_asc",
        "historico_config": policy.as_dict(),
        "cliente_pb": "CLIENTE X",
        "df_historial_full": hist,
    }
    config = build_config("diferencia_precio", ui)
    data = build_datos_exp("diferencia_precio", hist, config, ui)
    exp = ExpedienteComercial(
        nombre="Diferencia de precio",
        familia="precio",
        estrategia="PriceDifference",
        datos=data,
        contexto=PipelineContext(config=config),
    )
    exp.condiciones = [pd.DataFrame({"SKU": ["00123"], "PRECIO_BASE": [8.0]})]
    exp = Pipeline().ejecutar(exp)
    assert exp.resultado is not None

    # Pintar auditoría igual que view_handlers._ejecutar
    df_res = exp.resultado.dataframe
    nc_alertas = CreditNoteAuditor().auditar(
        hist,
        df_res,
        documentos_historial=config.get("documentos_historial"),
        modalidad="individual",
    )
    exp.resultado.metricas["nc_alertas"] = nc_alertas
    _df_x = exp.resultado.get_excel()
    CreditNoteAuditor.combinar_auditoria(df_res, nc_alertas)
    if _df_x is not df_res:
        CreditNoteAuditor.combinar_auditoria(_df_x, nc_alertas)
    return exp


def _run_do_pipeline(hist, policy):
    ui = {
        "modalidad": "individual",
        "historico_config": policy.as_dict(),
        "cliente_pb": "CLIENTE X",
        "df_historial_full": hist,
    }
    config = build_config("descuento_precio", ui)
    data = build_datos_exp("descuento_precio", hist, config, ui)
    exp = ExpedienteComercial(
        nombre="Descuento comercial",
        familia="precio",
        estrategia="PriceDifference",
        variante="discount_period",
        datos=data,
        contexto=PipelineContext(config=config),
    )
    exp.condiciones = [pd.DataFrame({"CODIGO": ["00123"], "DESCUENTO": [0.10]})]
    exp = Pipeline().ejecutar(exp)
    assert exp.resultado is not None

    # Pintar auditoría igual que view_handlers._ejecutar
    df_res = exp.resultado.dataframe
    nc_alertas = CreditNoteAuditor().auditar(
        hist,
        df_res,
        documentos_historial=config.get("documentos_historial"),
        modalidad="individual",
    )
    exp.resultado.metricas["nc_alertas"] = nc_alertas
    _df_x = exp.resultado.get_excel()
    CreditNoteAuditor.combinar_auditoria(df_res, nc_alertas)
    if _df_x is not df_res:
        CreditNoteAuditor.combinar_auditoria(_df_x, nc_alertas)
    return exp


def _sheet_values(path, sheet_prefer="Historico"):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_prefer] if sheet_prefer in wb.sheetnames else wb[wb.sheetnames[0]]
    vals = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c]
    wb.close()
    return vals


def test_expediente_carry_nc_rule(tmp_path):
    hist = _history()
    policy = HistorialConfig(facturas=(True, True), nc=(True, True), ndb=(True, False))
    exp = _run_pipeline(hist, policy)

    row = exp.resultado.dataframe.iloc[0]
    assert row["PRECIO_HIST"] == 9.0
    assert "APLICADA" in str(row["ALERTA"])
    assert "N900-1" in str(row["AUDITORIA_NC"])

    dirs = generar_expediente(
        resultado=exp,
        tipo_actual="DC",
        modalidad="individual",
        historico_config=policy,
        df_historial=hist,
        cliente_value="CLIENTE X",
        vendedor_value="",
        vendedor_display="",
        antecedentes="smoke",
        observaciones="",
        desktop_path=tmp_path,
    )
    assert len(dirs) == 1
    files = sorted(p.name for p in Path(dirs[0]).rglob("*") if p.is_file())
    assert len(files) == 3  # Calculo + Historico + Informe

    calc = next(Path(dirs[0]).rglob("*Calculo*.xlsx"))
    assert any("N900-1" in v for v in _sheet_values(calc, "NC_"))

    hist_xlsx = next(Path(dirs[0]).rglob("*Historico*.xlsx"))
    vals = _sheet_values(hist_xlsx)
    assert any("N900" in v for v in vals)  # NC visible (mostrar=True)
    assert any("001" in v and "10" in v for v in vals)  # factura con serie/nro


def test_expediente_historico_oculta_nc_por_mostrar(tmp_path):
    hist = _history()
    # NC usada en el cálculo (usar=True) pero NO visible en el Histórico
    policy = HistorialConfig(facturas=(True, True), nc=(False, True), ndb=(True, False))
    exp = _run_pipeline(hist, policy)

    dirs = generar_expediente(
        resultado=exp,
        tipo_actual="DC",
        modalidad="individual",
        historico_config=policy,
        df_historial=hist,
        cliente_value="CLIENTE X",
        vendedor_value="",
        vendedor_display="",
        antecedentes="smoke",
        observaciones="",
        desktop_path=tmp_path,
    )
    hist_xlsx = next(Path(dirs[0]).rglob("*Historico*.xlsx"))
    vals = _sheet_values(hist_xlsx)
    assert not any("N900" in v for v in vals)  # oculta pese a usar=True
    assert any("001" in v and "10" in v for v in vals)  # factura sí


def _all_values(path):
    wb = load_workbook(path, data_only=True)
    vals = [
        str(c)
        for name in wb.sheetnames
        for row in wb[name].iter_rows(values_only=True)
        for c in row
        if c
    ]
    wb.close()
    return vals


def test_expediente_do_carry_nc_rule(tmp_path):
    hist = _history()
    policy = HistorialConfig(facturas=(True, True), nc=(True, True), ndb=(True, False))
    exp = _run_do_pipeline(hist, policy)

    row = exp.resultado.dataframe.iloc[0]
    assert row["PRECIO_HIST"] == 9.0
    assert row["MONTO_NC"] == 90.0
    assert "APLICADA" in str(row["ALERTA"])
    assert "N900-1" in str(row["AUDITORIA_NC"])

    dirs = generar_expediente(
        resultado=exp,
        tipo_actual="DO",
        modalidad="individual",
        historico_config=policy,
        df_historial=hist,
        cliente_value="CLIENTE X",
        vendedor_value="",
        vendedor_display="",
        antecedentes="smoke",
        observaciones="",
        desktop_path=tmp_path,
    )
    assert len(dirs) == 1
    files = sorted(p.name for p in Path(dirs[0]).rglob("*") if p.is_file())
    assert len(files) == 3  # Calculo + Historico + Informe

    calc = next(Path(dirs[0]).rglob("*Calculo*.xlsx"))
    assert any("N900-1" in v for v in _all_values(calc))
    assert any("APLICADA" in v for v in _all_values(calc))

    hist_xlsx = next(Path(dirs[0]).rglob("*Historico*.xlsx"))
    vals = _sheet_values(hist_xlsx)
    assert any("N900" in v for v in vals)  # NC visible (mostrar=True)
    assert any("001" in v and "10" in v for v in vals)  # factura con serie/nro
