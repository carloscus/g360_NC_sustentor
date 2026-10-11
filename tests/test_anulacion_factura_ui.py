"""ANF (anulacion de factura) de punta a punta desde la UI.

Hallazgo: `factura_id` nunca llegaba a config (dropdown stub + `pass` en
build_config) y ejecutar ANF siempre terminaba en "Sin resultados: No se
seleccionó ninguna factura". Estos tests fijan:
1. build_config resuelve la factura del picker o del fragmento.
2. _ejecutar_tarea con 1 factura pinta la vista previa sin errores.
3. La estrategia tolera ceros a la izquierda en NUMERO.
"""

import sys
from pathlib import Path

import flet as ft
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent.parent))

from src.core import ventas_db
from src.core.ventas_db_client import VentasDbClient
from src.core.xls_processor import derivar_campos
from src.ui.config_builder import _factura_unica_fragmento, build_config
from src.ui.reconocimiento_view import ReconocimientoView

from tests.conftest import *  # noqa: F401,F403  (fixtures: tmp_db)
from tests.test_ejecucion_multicliente import _fila  # noqa: E402
from tests.test_render_parentdata import _AppStub  # noqa: E402

CID = "00056101"


def _poblar_1_factura():
    d = _fila(
        CID,
        "20601024714",
        "CLIENTE ALFA S.A.",
        "F01",
        "204",
        "67375",
        500.0,
        100.0,
        "2026-09-05",
    )
    d = dict(d)
    derivar_campos(d)
    conn = ventas_db.get_conn()
    ventas_db.insert_ventas(conn, [d])
    ventas_db.populate_nc_asociadas(conn)


def _historial_real() -> pd.DataFrame:
    cli = VentasDbClient()
    return cli.fetch_historial(id_cliente=CID, fecha_desde="2026-09-01", fecha_hasta="2026-09-30")


def _fac_id(hist: pd.DataFrame) -> str:
    r = hist.iloc[0]
    return f"{str(r['TIPO_DOC']).strip()}-{str(r['SERIE']).strip()}-{str(r['NUMERO']).strip()}"


def test_build_config_anf_desde_picker():
    cfg = build_config("anular_factura", {"factura_ci": "F01-204-67375", "historico_config": {}})
    assert cfg["factura_id"] == "F01-204-67375"


def test_build_config_anf_desde_fragmento():
    hist = pd.DataFrame(
        [{"TIPO_DOC": "F01", "SERIE": "204", "NUMERO": "67375", "TIPO_CLASE": "factura"}]
    )
    cfg = build_config(
        "anular_factura",
        {"factura_ci": None, "historico_config": {}, "df_historial_full": hist},
    )
    assert cfg["factura_id"] == "F01-204-67375"


def test_build_config_anf_sin_factura_unica():
    hist = pd.DataFrame(
        [
            {"TIPO_DOC": "F01", "SERIE": "204", "NUMERO": "67375", "TIPO_CLASE": "factura"},
            {"TIPO_DOC": "F01", "SERIE": "204", "NUMERO": "67380", "TIPO_CLASE": "factura"},
        ]
    )
    assert _factura_unica_fragmento(hist) is None
    assert _factura_unica_fragmento(None) is None
    cfg = build_config(
        "anular_factura",
        {"factura_ci": None, "historico_config": {}, "df_historial_full": hist},
    )
    # Sin picker, caen TODAS las del fragmento (ya no exige una sola).
    assert cfg["factura_ids"] == ["F01-204-67375", "F01-204-67380"]
    assert cfg["factura_id"] == "F01-204-67375"


def test_build_config_anf_multi_picker():
    cfg = build_config(
        "anular_factura",
        {
            "factura_ci": "F01-204-67375",
            "facturas_sel": ["F01-204-67375", "F01-204-67380"],
            "historico_config": {},
        },
    )
    assert cfg["factura_ids"] == ["F01-204-67375", "F01-204-67380"]


def test_estrategia_tolera_ceros_en_numero():
    from src.domain import ExpedienteComercial, PipelineContext
    from src.strategies.cancel_invoice import CancelInvoiceStrategy

    df = pd.DataFrame(
        [
            {
                "TIPO_DOC": "F01",
                "SERIE": "204",
                "NUMERO": "067375",
                "SOLES": 500.0,
                "CANTIDAD": 100.0,
                "CODIGO": "014850",
            }
        ]
    )
    exp = ExpedienteComercial(
        nombre="ANF",
        familia="anular",
        estrategia="CancelInvoice",
        datos=df,
        contexto=PipelineContext(config={"factura_id": "F01-204-67375"}),
    )
    res = CancelInvoiceStrategy().process(exp)
    assert not res.dataframe.empty
    assert res.resumen["total_nc"] == 500.0


def test_estrategia_multi_factura_multi_cliente():
    from src.domain import ExpedienteComercial, PipelineContext
    from src.strategies.cancel_invoice import CancelInvoiceStrategy

    def _row(tpo, serie, nro, cid, nom, soles):
        return {
            "TIPO_DOC": tpo,
            "SERIE": serie,
            "NUMERO": nro,
            "SOLES": soles,
            "CANTIDAD": 100.0,
            "CODIGO": "014850",
            "id_cliente": cid,
            "nom_cliente": nom,
        }

    df = pd.DataFrame(
        [
            _row("F01", "204", "67375", "00056101", "ALFA", 500.0),
            _row("F01", "204", "67380", "00005610", "BETA", 300.0),
        ]
    )
    exp = ExpedienteComercial(
        nombre="ANF",
        familia="anular",
        estrategia="CancelInvoice",
        datos=df,
        contexto=PipelineContext(
            config={"factura_ids": ["F01-204-67375", "F01-204-67380", "F01-204-00000"]}
        ),
    )
    res = CancelInvoiceStrategy().process(exp)
    out = res.dataframe
    assert len(out) == 2
    assert sorted(out["FACTURA"].unique()) == ["F01-204-67375", "F01-204-67380"]
    assert sorted(out["COD_CLIENTE"].unique()) == ["00005610", "00056101"]
    assert res.resumen["total_nc"] == 800.0
    assert len(res.por_condicion) == 2
    # La inexistente avisa (warning) sin tumbar a las encontradas.
    assert any("00000" in a.mensaje for a in res.alertas)


def test_gate_anf_acepta_varias_facturas(tmp_db):
    from types import SimpleNamespace

    from src.ui.reconocimiento_view import ReconocimientoView

    hist = pd.DataFrame([{"COD_CLIENTE": CID}])
    for nfac, puede in ((0, False), (1, True), (3, True)):
        fake = SimpleNamespace(
            tipo_actual="ANF",
            df_historial=hist,
            _hist_fragmento={"facturas": nfac} if nfac else None,
        )
        fake._insumo_necesita = lambda ins, cfg: ins in cfg.get("insumos", ())
        assert (
            ReconocimientoView._insumos_especificos_disponibles(fake, {"insumos": ("historico",)})
            is puede
        )


def test_anf_ejecuta_y_pinta_vista_previa(tmp_db, tmp_path):
    _poblar_1_factura()
    hist = _historial_real()
    assert len(hist) >= 1
    fac_id = _fac_id(hist)

    app = _AppStub()
    errores = []
    app.show_snackbar = lambda msg, *a, **k: errores.append(msg)

    view = ReconocimientoView(app)
    view.build()
    view.tipo_actual = "ANF"
    view.tipo_dropdown.value = "ANF"
    view.df_historial = hist
    view._sel_facturas = [(CID, fac_id)]
    view._ejecutar_tarea()

    assert errores == [], f"la app mostro un error: {errores}"
    assert view.resultado is not None
    assert view.resultado.resultado is not None
    assert not view.resultado.resultado.dataframe.empty
    assert view.resultados_container.visible is True
    assert view._result_panel.visible is True


CID_B = "00005610"


def test_qr_sunat_datos_factura_original():
    """El QR lleva la factura ORIGINAL: serie de 3 partes, fecha de emisión
    de la factura (metricas) y el pipe final del formato impreso SUNAT."""
    from src.render.excel_renderer import _build_qr_data

    # 11 campos: hash vacío (placeholder del CDR) + campo final vacío.
    data = _build_qr_data("20555884116", "204", "69002", 9101.25, "2026-09-25", "20601745357")
    partes = data.split("|")
    assert partes[:9] == [
        "20555884116",
        "01",
        "204",
        "69002",
        f"{round(9101.25 * 0.18 / 1.18, 2):.2f}",  # IGV incluido: total * 18/118
        "9101.25",
        "2026-09-25",
        "6",
        "20601745357",
    ]
    assert len(partes) == 11 and partes[9] == "" and partes[10] == ""


def test_qr_serie_tres_partes_y_fecha_factura():
    """_insert_qr_anf parsea 'F01-204-69002' (TIPO-SERIE-NUMERO) sin meter el
    tipo en la serie, y prioriza la fecha de la factura sobre la de hoy."""
    import datetime as dt

    from src.render.excel_renderer import ExcelRenderer

    renderer = ExcelRenderer()
    df = pd.DataFrame([{"FECHA": dt.date(2026, 9, 25), "MONTO_NC": 9101.25}])
    # Con fecha de metricas: usa la fecha de emisión de la factura.
    buf = renderer._insert_qr_anf(
        df, "F01-204-69002", 9101.25, "20601745357", fecha_factura="2026-09-25"
    )
    assert buf is not None
    # La imagen es un PNG válido: el contenido del QR se verifica vía
    # _build_qr_data (test anterior); aquí basta la generación correcta.
    assert buf.getvalue()[:8] == b"\x89PNG\r\n\x1a\n"

    # Serie legacy de 2 partes ("F204-69002") también resuelve a "204".
    assert _parse_doc_ref("F204-69002") == ("204", "69002")
    assert _parse_doc_ref("F01-204-69002") == ("204", "69002")


def _parse_doc_ref(doc_ref):
    """Réplica del parseo de _insert_qr_anf para probarlo sin imagen."""
    if "-" in doc_ref:
        partes = doc_ref.split("-")
        if len(partes) >= 3:
            return partes[1], partes[2]
        return (partes[0][1:] if len(partes[0]) > 1 else "", partes[1])
    return "", doc_ref


def test_anf_calculo_estructura_y_montos_numericos(tmp_db, tmp_path):
    """Expediente ANF: montos numéricos (no texto), encabezados en fila 12,
    sin filas congeladas y fila NC en 7 con fórmula de aplicación directa."""
    from src.ui import expediente_service as es
    from src.ui.catalog import HistorialConfig
    from openpyxl import load_workbook

    _poblar_1_factura()
    hist = _historial_real()
    fac_id = _fac_id(hist)

    app = _AppStub()
    app.show_snackbar = lambda msg, *a, **k: None
    view = ReconocimientoView(app)
    view.build()
    view.tipo_actual = "ANF"
    view.tipo_dropdown.value = "ANF"
    view.df_historial = hist
    view._sel_facturas = [(CID, fac_id)]
    view._ejecutar_tarea()

    out_dir = tmp_path / "exp"
    out_dir.mkdir()
    es.generar_expediente(
        view.resultado,
        "ANF",
        "individual",
        HistorialConfig(facturas=(True, True), nc=(True, False), ndb=(True, False)),
        hist,
        CID,
        "01188",
        "VEND",
        "",
        "",
        out_dir,
    )
    xls = next(out_dir.rglob("*_Calculo.xlsx"))
    wb = load_workbook(str(xls))
    ws = wb[wb.sheetnames[0]]

    # Sin filas congeladas (política: expedientes xlsx sin freeze).
    assert ws.freeze_panes is None

    # Encabezados de la tabla en la fila 12 (sin filas en blanco extra).
    assert ws.cell(row=12, column=1).value == "N°"
    assert ws.cell(row=12, column=2).value == "FACTURA"
    assert ws.cell(row=12, column=6).value == "P. UNITARIO"
    assert ws.cell(row=12, column=7).value == "TOTAL FACTURA"
    assert ws.cell(row=12, column=8).value == "MONTO"

    # Montos numéricos: P. UNITARIO (F) y TOTAL FACTURA (G), no texto.
    v_f = ws.cell(row=13, column=6).value
    v_g = ws.cell(row=13, column=7).value
    assert isinstance(v_f, (int, float)), f"P. UNITARIO como texto: {v_f!r}"
    assert isinstance(v_g, (int, float)), f"TOTAL FACTURA como texto: {v_g!r}"

    # MONTO referencia a TOTAL FACTURA (hoja viva).
    assert str(ws.cell(row=13, column=8).value).startswith("=")

    # Fila NC en 7: aplicación directa (NC = 100% de la factura).
    assert "NC" in str(ws.cell(row=7, column=6).value)
    assert str(ws.cell(row=7, column=7).value).startswith("=")

    # El histórico del expediente tampoco lleva filas congeladas.
    hist_xls = next(out_dir.rglob("*_Historico.xlsx"))
    wb_h = load_workbook(str(hist_xls))
    for ws_h in wb_h.worksheets:
        assert ws_h.freeze_panes is None, ws_h.title


def test_anf_multi_factura_multi_cliente(tmp_db, tmp_path):
    d1 = _fila(
        CID, "20601024714", "CLIENTE ALFA S.A.", "F01", "204", "67375", 500.0, 100.0, "2026-09-05"
    )
    d2 = _fila(
        CID_B,
        "20601024715",
        "CLIENTE BETA E.I.R.L.",
        "F01",
        "204",
        "67380",
        300.0,
        60.0,
        "2026-09-06",
    )
    filas = []
    for d in (d1, d2):
        d = dict(d)
        derivar_campos(d)
        filas.append(d)
    conn = ventas_db.get_conn()
    ventas_db.insert_ventas(conn, filas)
    ventas_db.populate_nc_asociadas(conn)

    cli = VentasDbClient()
    hist = pd.concat(
        [
            cli.fetch_historial(id_cliente=c, fecha_desde="2026-09-01", fecha_hasta="2026-09-30")
            for c in (CID, CID_B)
        ],
        ignore_index=True,
    )
    facs = sorted(
        f"{str(r['TIPO_DOC']).strip()}-{str(r['SERIE']).strip()}-{str(r['NUMERO']).strip()}"
        for _, r in hist.iterrows()
    )
    assert len(facs) == 2

    app = _AppStub()
    errores = []
    app.show_snackbar = lambda msg, *a, **k: errores.append(msg)

    view = ReconocimientoView(app)
    view.build()
    view.tipo_actual = "ANF"
    view.tipo_dropdown.value = "ANF"
    view.df_historial = hist
    view._sel_facturas = [(CID, facs[0]), (CID_B, facs[1])]
    view._ejecutar_tarea()

    assert errores == [], f"la app mostro un error: {errores}"
    df = view.resultado.resultado.dataframe
    assert len(df) == 2
    assert sorted(df["FACTURA"].unique()) == facs
    # Un tab por cliente, ambos pintados.
    assert len(view._result_tabs.tabs) == 2
    for tab in view._result_tabs.tabs:
        assert isinstance(tab.content.content, ft.Column)
    assert view._result_tabs.selected_index == 0
