"""CF de seguimiento en la hoja Cálculo: DIF. UNITARIA y precio histórico.

DIF. UNITARIA > 0 en verde (filas que generan NC), = 0 en gris (se apagan);
precio histórico = 0 en ámbar (sin base no hay con qué comparar). Solo
reglas CellIs sobre el valor (era-2007: las evalúan Excel y LibreOffice).
"""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.render.excel_renderer import ExcelRenderer
from tests.test_render_tipos_smoke import TOTALES, _df, _res


def _generar(tmp_path, tipo):
    out = tmp_path / f"{tipo}_cf.xlsx"
    ExcelRenderer().generar(
        _res(_df(tipo), tipo, TOTALES[tipo]),
        str(out),
        tipo=tipo,
        cliente="CLIENTE X",
        doc_ref="F201-100",
    )
    wb = load_workbook(str(out))
    return wb[wb.sheetnames[0]]


def _reglas(ws):
    """{rango: [operadores]} de todas las reglas condicionales."""
    out = {}
    for cf in ws.conditional_formatting:
        out.setdefault(str(cf.sqref), []).extend(rl.operator for rl in cf.rules)
    return out


def _ops(reglas, col, fila):
    """Operadores sobre una celda (openpyxl normaliza F20:F20 → F20)."""
    return reglas.get(f"{col}{fila}:{col}{fila}", reglas.get(f"{col}{fila}", []))


def _fila_encabezado(ws, texto):
    for row in ws.iter_rows():
        if texto in [c.value for c in row]:
            return row[0].row, [c.value for c in row]
    raise AssertionError(f"sin fila de encabezado con {texto!r}")


class TestCfDual:
    """diferencia_precio: tabla 1 (PRECIO UNID.) + tabla 2 (HIST + DIF)."""

    def test_dif_unitaria_verde_y_gris(self, tmp_path):
        ws = _generar(tmp_path, "diferencia_precio")
        r_hdr, headers = _fila_encabezado(ws, "DIF. UNITARIA")
        col = get_column_letter(headers.index("DIF. UNITARIA") + 1)
        reglas = _reglas(ws)
        ops = _ops(reglas, col, r_hdr + 1)
        assert "greaterThan" in ops, "DIF > 0 debe ir en verde"
        assert "equal" in ops, "DIF = 0 debe apagarse en gris"

    def test_precio_hist_efectivo_ambar(self, tmp_path):
        ws = _generar(tmp_path, "diferencia_precio")
        r_hdr, headers = _fila_encabezado(ws, "PRECIO HIST. EFECTIVO")
        col = get_column_letter(headers.index("PRECIO HIST. EFECTIVO") + 1)
        reglas = _reglas(ws)
        assert "equal" in _ops(reglas, col, r_hdr + 1)

    def test_precio_unid_tabla1_ambar(self, tmp_path):
        """El histórico de tabla 1 (COMO SE ATENDIÓ) también se marca."""
        ws = _generar(tmp_path, "diferencia_precio")
        r_hdr, headers = _fila_encabezado(ws, "PRECIO UNID.")
        col = get_column_letter(headers.index("PRECIO UNID.") + 1)
        reglas = _reglas(ws)
        assert "equal" in _ops(reglas, col, r_hdr + 1)


class TestCfSimple:
    """descuento_precio: tabla única (PRECIO ATENDIDO + DIF. UNITARIA)."""

    def test_dif_y_atendido(self, tmp_path):
        ws = _generar(tmp_path, "descuento_precio")
        r_hdr, headers = _fila_encabezado(ws, "DIF. UNITARIA")
        col_dif = get_column_letter(headers.index("DIF. UNITARIA") + 1)
        col_hist = get_column_letter(headers.index("PRECIO ATENDIDO") + 1)
        reglas = _reglas(ws)
        ops_dif = _ops(reglas, col_dif, r_hdr + 1)
        assert "greaterThan" in ops_dif and "equal" in ops_dif
        assert "equal" in _ops(reglas, col_hist, r_hdr + 1)
