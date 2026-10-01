"""Diseño del Informe de Sustento (orden, tablas y nombres reales)."""

import docx
import pandas as pd

from src.domain import RecognitionResult
from src.render.docx_renderer import DocxInformeRenderer


def _df_dc():
    return pd.DataFrame(
        [
            {
                "SKU": "02203",
                "ARTICULO": "FORRO N VINIFAN A4 CRISTAL 25",
                "CANTIDAD": 4000,
                "SOLES": 13878.44,
                "PRECIO_HIST": 3.46961,
                "PRECIO_BASE": 5.12,
                "PRECIO_NETO": 3.33082,
                "DIFERENCIA": 0.13879,
                "MONTO_NC": 555.16,
                "FACTURA": "F204-40260",
                "ALERTA": "AL01 - Diferencia positiva (SKU 02203)",
            },
            {
                "SKU": "02202",
                "ARTICULO": "FORRO N VINIFAN OFICIO CRISTAL 25",
                "CANTIDAD": 2500,
                "SOLES": 17364.99,
                "PRECIO_HIST": 6.946,
                "PRECIO_BASE": 10.25,
                "PRECIO_NETO": 6.66816,
                "DIFERENCIA": 0.27784,
                "MONTO_NC": 694.60,
                "FACTURA": "F204-40260",
                "ALERTA": "AL01 - Diferencia positiva (SKU 02202)",
            },
        ]
    )


def _generar(tmp_path, df=None, **datos_extra):
    df = _df_dc() if df is None else df
    total = round(float(df["MONTO_NC"].sum()), 2)
    res = RecognitionResult(
        dataframe=df,
        resumen={"total_nc": total, "skus_afectados": 2, "doc_ref": "F204-40260"},
    )
    datos = {
        "tipo_operacion": "DC",
        "cliente": "CLIENTE X",
        "representante": "V",
        "numero_referencia": "F204-40260",
        "evidencias": {},
        "modalidad": "individual",
        "periodo": "11/09/2025 al 25/04/2026",
        "descripcion": "descuentos omitidos segun lista de precios",
        "observaciones": "descuentos omitidos en facturas",
    }
    datos.update(datos_extra)
    out = tmp_path / "EXP-DC-1_Informe.docx"
    DocxInformeRenderer().generar(
        res,
        ruta_salida=str(out),
        datos_adicionales=datos,
        nombre_archivo_excel="EXP-DC-1_Calculo.xlsx",
    )
    return docx.Document(str(out))


def _parrafos(doc):
    return [p.text for p in doc.paragraphs if p.text.strip()]


def _tablas_txt(doc):
    return [[c.text for c in r.cells] for t in doc.tables for r in t.rows]


class TestOrden:
    def test_resultado_antes_que_analisis(self, tmp_path):
        pars = _parrafos(_generar(tmp_path))
        i_res = next(i for i, t in enumerate(pars) if "RESULTADO ECON" in t)
        i_ana = next(i for i, t in enumerate(pars) if "ANTECEDENTES" in t)
        i_obs = next(i for i, t in enumerate(pars) if "OBSERVACIONES" in t)
        i_doc = next(i for i, t in enumerate(pars) if t == "5. DOCUMENTOS")
        assert i_res < i_ana < i_obs < i_doc
        assert not any("PROCESO APLICADO" in t for t in pars)


class TestDatosGenerales:
    def test_filas_clave(self, tmp_path):
        rows = _tablas_txt(_generar(tmp_path))
        plano = [c for r in rows for c in r]
        assert "Modalidad" not in plano
        assert "Doc. de referencia" in plano
        assert "F204-40260" in plano
        assert "Alcance" in plano
        assert "2 SKU · 6,500 unidades" in plano

    def test_labels_laterales_en_10pt(self, tmp_path):
        from docx.shared import Pt

        doc = _generar(tmp_path)
        label_runs = [
            r
            for t in doc.tables
            for row in t.rows
            for r in row.cells[0].paragraphs[0].runs
            if r.bold
        ]
        assert label_runs
        assert all(r.font.size == Pt(10) for r in label_runs)

    def test_margenes_laterales_simetricos_2cm(self, tmp_path):
        doc = _generar(tmp_path)
        sec = doc.sections[0]
        assert abs(sec.left_margin.cm - 2.0) < 0.01
        assert abs(sec.right_margin.cm - 2.0) < 0.01
        assert abs(sec.left_margin.cm - sec.right_margin.cm) < 0.001

    def test_sin_fila_duplicada_de_facturas(self, tmp_path):
        rows = _tablas_txt(_generar(tmp_path))
        plano = [c for r in rows for c in r]
        assert "Facturas comprometidas" not in plano
        assert plano.count("F204-40260") == 1

    def test_consolidado_muestra_lista(self, tmp_path):
        df = _df_dc().copy()
        df.loc[1, "FACTURA"] = "F204-40261"
        doc = _generar(tmp_path, df=df, modalidad="consolidado")
        plano = [c for r in _tablas_txt(doc) for c in r]
        i = plano.index("Doc. de referencia")
        assert plano[i + 1] == "F204-40260, F204-40261"


class TestSinDatosDeCalculo:
    def test_sin_tabla_detalle_ni_montos_por_sku(self, tmp_path):
        doc = _generar(tmp_path)
        rows = _tablas_txt(doc)
        plano = [c for r in rows for c in r]
        # El informe explica; el detalle vive en el Cálculo.
        assert "P.U. fact." not in plano
        assert "02203" not in plano
        assert "3.46961" not in plano
        assert "S/ 555.16" not in plano
        pars = _parrafos(doc)
        assert not any("SKU afectados" in t for t in pars)
        # La conclusión en dinero sí va (tabla de 3 filas).
        assert "S/ 1,249.76" in plano


class TestNormalizacion:
    def test_mayuscula_y_punto(self, tmp_path):
        pars = _parrafos(_generar(tmp_path))
        assert "Descuentos omitidos según lista de precios." in pars
        assert "Descuentos omitidos en facturas." in pars

    def test_corrige_typos_frecuentes(self, tmp_path):
        pars = _parrafos(
            _generar(
                tmp_path,
                descripcion="diferencia de prercios - atencion con descuentos omitidos segun lista",
                observaciones="atencion con omision de descuento",
            )
        )
        assert "Diferencia de precios - atención con descuentos omitidos según lista." in pars
        assert "Atención con omisión de descuento." in pars


class TestDocumentos:
    def _bullets(self, tmp_path, **kw):
        doc = _generar(tmp_path, **kw)
        return [p.text for p in doc.paragraphs if p.text.strip()]

    def test_lista_con_nombres_reales(self, tmp_path):
        pars = self._bullets(tmp_path)
        assert "EXP-DC-1_Calculo.xlsx — Hoja de cálculo" in pars
        assert "EXP-DC-1_Informe.docx — Este informe" in pars
        assert not any("Informe de Sustento Comercial.docx" in t for t in pars)
        assert not any(t.startswith("Archivos:") for t in pars)

    def test_historico_cuando_se_informa(self, tmp_path):
        pars = self._bullets(tmp_path, nombre_historico="EXP-DC-1_Historico.xlsx")
        assert ("EXP-DC-1_Historico.xlsx — Segmento histórico del ERP") in pars


class TestFraseCierre:
    def test_dc_individual(self, tmp_path):
        pars = _parrafos(_generar(tmp_path))
        assert (
            "Se revisó el histórico frente a la lista de precios vigente; "
            "la nota de crédito por S/ 1,249.76 regulariza las diferencias "
            "en 2 SKU de la factura F204-40260."
        ) in pars

    def test_dc_consolidado(self, tmp_path):
        df = _df_dc().copy()
        df.loc[1, "FACTURA"] = "F204-40261"
        pars = _parrafos(_generar(tmp_path, df=df, modalidad="consolidado"))
        assert (
            "Se revisó el histórico frente a la lista de precios vigente; "
            "la nota de crédito por S/ 1,249.76 regulariza las diferencias "
            "en 2 SKU de 2 factura(s)."
        ) in pars

    def test_descuento(self, tmp_path):
        pars = _parrafos(_generar(tmp_path, tipo_calculo="descuento_precio"))
        assert (
            "Se revisaron los descuentos omitidos frente al histórico; "
            "la nota de crédito por S/ 1,249.76 regulariza lo no aplicado "
            "en 2 SKU."
        ) in pars

    def test_otro_tipo_fallback_con_importe(self, tmp_path):
        pars = _parrafos(_generar(tmp_path, tipo_calculo="devolucion_fisica"))
        assert any("Importe regularizado: S/ 1,249.76 en 2 SKU." in t for t in pars)

    def test_anular(self, tmp_path):
        pars = _parrafos(_generar(tmp_path, tipo_calculo="anular_factura"))
        assert any(
            "La factura F204-40260 se anula al 100%; la nota de crédito "
            "por S/ 1,249.76 regulariza el documento." in t
            for t in pars
        )

    def test_rebate(self, tmp_path):
        df = _df_dc().copy()
        df["LINEA"] = "L01"
        pars = _parrafos(_generar(tmp_path, df=df, tipo_calculo="rebate_volumen"))
        assert any(
            "la nota de crédito por S/ 1,249.76 regulariza el rebate en línea L01." in t
            for t in pars
        )

    def test_bonificacion(self, tmp_path):
        pars = _parrafos(_generar(tmp_path, tipo_calculo="bonificacion_promocion"))
        assert any(
            "la nota de crédito por S/ 1,249.76 regulariza la bonificación en 2 SKU." in t
            for t in pars
        )

    def test_feria(self, tmp_path):
        pars = _parrafos(_generar(tmp_path, tipo_calculo="feria_preventa"))
        assert any(
            "la nota de crédito por S/ 1,249.76 regulariza lo comprometido en 2 SKU." in t
            for t in pars
        )
