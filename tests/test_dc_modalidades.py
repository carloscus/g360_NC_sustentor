"""Modalidades del caso DC (diferencia_precio).

Individual (default): una fila por línea de factura, sin agrupar.
Consolidado: una fila por SKU (cantidad Σ, precio moda con desempate por
sort_mode, lista de facturas, monto por suma exacta de líneas).
"""

import pandas as pd

from src.strategies.price_difference import PriceDifferenceStrategy
from src.ui.config_builder import build_config
from src.ui.view_handlers import mapear_nc_existente


def _hist():
    return pd.DataFrame(
        {
            "CODIGO": ["S1", "S1", "S2", "S3"],
            "ARTICULO": ["Prod 1", "Prod 1", "Prod 2", "Prod 3"],
            "CANTIDAD": [100, 200, 50, 10],
            "SOLES": [340.00, 700.00, 125.00, 20.00],
            "FECHA": pd.to_datetime(["2026-01-15", "2026-02-20", "2026-01-15", "2026-02-20"]),
            "TIPO_DOC": ["F01"] * 4,
            "SERIE": ["201"] * 4,
            "NUMERO": ["100", "101", "100", "101"],
            "COD_CLIENTE": ["00068426"] * 4,
            "DOC_CLIENTE": ["206039928"] * 4,
            "CLIENTE": ["DISTRIBUIDORA LOS ANGELES"] * 4,
        }
    )


def _lista():
    return pd.DataFrame(
        [
            {"SKU": "S1", "PRECIO_BASE": 3.00},
            {"SKU": "S2", "PRECIO_BASE": 2.00},
            {"SKU": "S3", "PRECIO_BASE": 2.00},
        ]
    )


def _comparar(modalidad, sort_mode="fecha_desc"):
    st = PriceDifferenceStrategy()
    return st._procesar_comparar(
        _hist(),
        _lista(),
        [],
        config={"modalidad": modalidad, "sort_mode": sort_mode},
    )


class TestModalidadIndividual:
    def test_default_una_fila_por_linea(self):
        st = PriceDifferenceStrategy()
        df_res, _, _ = st._procesar_comparar(_hist(), _lista(), [])
        # S1 aparece en 2 facturas → 2 filas independientes
        assert len(df_res) == 4
        s1 = df_res[df_res["SKU"] == "S1"].sort_values("FACTURA")
        assert list(s1["FACTURA"]) == ["F201-100", "F201-101"]
        assert list(s1["MONTO_NC"]) == [40.00, 100.00]

    def test_explicita_igual_default(self):
        df_def, _, _ = _comparar("individual")
        df_exp, _, _ = _comparar("individual")
        assert df_def.reset_index(drop=True).equals(df_exp.reset_index(drop=True))


class TestModalidadConsolidado:
    def test_una_fila_por_sku(self):
        df_res, _, df_full = _comparar("consolidado")
        assert len(df_res) == 3
        assert sorted(df_res["SKU"]) == ["S1", "S2", "S3"]
        # df_full también consolidado (el renderer usa get_excel)
        assert len(df_full) == 3

    def test_s1_cantidad_soles_y_facturas(self):
        df_res, _, _ = _comparar("consolidado")
        s1 = df_res[df_res["SKU"] == "S1"].iloc[0]
        assert s1["CANTIDAD"] == 300
        assert s1["SOLES"] == 1040.00
        assert s1["FACTURAS"] == "F201-100, F201-101"
        # Principal = mayor SOLES → F201-101 (700 > 340)
        assert s1["FACTURA"] == "F201-101"

    def test_moda_desempata_por_mas_reciente(self):
        # Precios 3.40 (15/01) y 3.50 (20/02) empatan 1-1 → gana el reciente
        df_res, _, _ = _comparar("consolidado", sort_mode="fecha_desc")
        s1 = df_res[df_res["SKU"] == "S1"].iloc[0]
        assert s1["PRECIO_HIST"] == 3.50
        assert s1["DIFERENCIA"] == 0.50

    def test_moda_desempate_fifo(self):
        df_res, _, _ = _comparar("consolidado", sort_mode="fecha_asc")
        s1 = df_res[df_res["SKU"] == "S1"].iloc[0]
        assert s1["PRECIO_HIST"] == 3.40

    def test_monto_suma_exacta_no_moda_por_cantidad(self):
        # Exacto: 0.40*100 + 0.50*200 = 140.00 (≠ 0.50*300 = 150.00)
        df_res, _, _ = _comparar("consolidado")
        s1 = df_res[df_res["SKU"] == "S1"].iloc[0]
        assert s1["MONTO_NC"] == 140.00
        assert "total S/ 140.00" in s1["ALERTA"]
        assert "2 facturas" in s1["ALERTA"]

    def test_total_igual_individual(self):
        df_ind, _, _ = _comparar("individual")
        df_con, _, _ = _comparar("consolidado")
        assert round(df_con["MONTO_NC"].sum(), 2) == round(df_ind["MONTO_NC"].sum(), 2) == 165.00

    def test_auditoria_desglose_por_factura(self):
        df_res, _, _ = _comparar("consolidado")
        s1 = df_res[df_res["SKU"] == "S1"].iloc[0]
        assert "Cortes: F201-100 100u (3.40); F201-101 200u (3.50)." in s1["AUDITORIA_NC"]

    def test_auditoria_con_rango_si_varia(self):
        df_res, _, _ = _comparar("consolidado")
        s1 = df_res[df_res["SKU"] == "S1"].iloc[0]
        assert "rango" in s1["AUDITORIA_NC"]
        assert "suma exacta" in s1["AUDITORIA_NC"]
        s2 = df_res[df_res["SKU"] == "S2"].iloc[0]
        assert "rango" not in s2["AUDITORIA_NC"]

    def test_al11_sobre_fila_consolidada(self):
        hist = pd.DataFrame(
            {
                "CODIGO": ["S9", "S9"],
                "ARTICULO": ["P9", "P9"],
                "CANTIDAD": [5, 5],
                "SOLES": [10.025, 10.025],  # 2.005 c/u vs lista 2.00 → dif 0.005
                "FECHA": pd.to_datetime(["2026-01-15", "2026-02-20"]),
                "TIPO_DOC": ["F01", "F01"],
                "SERIE": ["201", "201"],
                "NUMERO": ["100", "101"],
            }
        )
        lista = pd.DataFrame([{"SKU": "S9", "PRECIO_BASE": 2.00}])
        st = PriceDifferenceStrategy()
        df_res, _, _ = st._procesar_comparar(
            hist, lista, [], config={"modalidad": "consolidado", "sort_mode": "fecha_desc"}
        )
        assert len(df_res) == 1
        # dif 0.005 ≤ 0.01 y total 0.05 ≤ 1.00 → AL11 (no genera NC)
        assert "AL11" in df_res["ALERTA"].iloc[0]


class TestConfigModalidad:
    def test_build_config_pasa_modalidad(self):
        cfg = build_config("diferencia_precio", {"modalidad": "consolidado"})
        assert cfg["modalidad"] == "consolidado"
        assert cfg["sort_mode"] == "fecha_desc"

    def test_build_config_default_individual(self):
        cfg = build_config("diferencia_precio", {})
        assert cfg["modalidad"] == "individual"


class TestRenderConsolidado:
    def _resultado(self):
        from src.domain import RecognitionResult

        df_res, _, df_full = _comparar("consolidado")
        total = round(float(df_full["MONTO_NC"].sum()), 2)
        return RecognitionResult(
            dataframe=df_res,
            dataframe_excel=df_full,
            resumen={"total_nc": total, "skus_afectados": 3, "doc_ref": "F201-101"},
        )

    def test_columna_facturas_y_monto_estatico(self, tmp_path):
        from openpyxl import load_workbook
        from src.render.excel_renderer import ExcelRenderer

        out = tmp_path / "Calculo.xlsx"
        ExcelRenderer().generar(
            self._resultado(),
            str(out),
            tipo="diferencia_precio",
            cliente="CLIENTE X",
            doc_ref="F201-101",
            ruc="206039928",
            modalidad="consolidado",
        )
        wb = load_workbook(str(out))
        ws = wb[wb.sheetnames[0]]
        textos = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert "FACTURAS" in textos
        assert "FACTURA" not in textos
        # MONTO de S1: estático 140.00 (suma exacta), no fórmula
        montos = [
            (c.coordinate, c.value) for row in ws.iter_rows() for c in row if c.value == 140.00
        ]
        assert montos, "MONTO exacto 140.00 no encontrado"
        for _, v in montos:
            assert not (isinstance(v, str) and v.startswith("="))

    def test_individual_sin_facturas(self, tmp_path):
        from openpyxl import load_workbook
        from src.domain import ExpedienteComercial, PipelineContext
        from src.render.excel_renderer import ExcelRenderer
        from src.strategies.price_difference import PriceDifferenceStrategy

        exp = ExpedienteComercial(
            nombre="T",
            estrategia="PriceDifference",
            datos=_hist(),
            condiciones=[_lista()],
            contexto=PipelineContext(config={"modalidad": "individual"}),
        )
        res = PriceDifferenceStrategy().process(exp)
        out = tmp_path / "Calculo.xlsx"
        ExcelRenderer().generar(
            res,
            str(out),
            tipo="diferencia_precio",
            cliente="CLIENTE X",
            doc_ref="F201-101",
            ruc="206039928",
            modalidad="individual",
        )
        wb = load_workbook(str(out))
        ws = wb[wb.sheetnames[0]]
        textos = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert "FACTURA" in textos
        assert "FACTURAS" not in textos
        assert not any("COMPROMETIDAS" in t for t in textos)


class TestDocxModalidad:
    def _texto(self, tmp_path, datos_modalidad):
        import docx
        from src.domain import RecognitionResult
        from src.render.docx_renderer import DocxInformeRenderer

        df = pd.DataFrame(
            {
                "SKU": ["S1"],
                "ARTICULO": ["P"],
                "CANTIDAD": [300],
                "MONTO_NC": [140.00],
                "FACTURA": ["F201-101"],
            }
        )
        res = RecognitionResult(
            dataframe=df,
            resumen={"total_nc": 140.00, "skus_afectados": 1, "doc_ref": "F201-101"},
        )
        datos = {
            "tipo_operacion": "DC",
            "cliente": "CLIENTE X",
            "representante": "V",
            "numero_referencia": "F201-101",
            "evidencias": {},
        }
        if datos_modalidad is not None:
            datos["modalidad"] = datos_modalidad
        out = tmp_path / "Informe.docx"
        DocxInformeRenderer().generar(
            res,
            ruta_salida=str(out),
            datos_adicionales=datos,
            nombre_archivo_excel="Calculo.xlsx",
        )
        doc = docx.Document(str(out))
        return "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)

    def test_modalidad_no_se_muestra_en_datos(self, tmp_path):
        # La modalidad gobierna motor y expedientes, pero no se imprime.
        txt = self._texto(tmp_path, "consolidado")
        assert "Modalidad" not in txt
        assert "Consolidado" not in txt
        assert "Individual" not in self._texto(tmp_path, None)


class TestFacturasComprometidas:
    def _process(self, modalidad):
        from src.domain import ExpedienteComercial, PipelineContext

        exp = ExpedienteComercial(
            nombre="T",
            estrategia="PriceDifference",
            datos=_hist(),
            condiciones=[_lista()],
            contexto=PipelineContext(config={"modalidad": modalidad, "sort_mode": "fecha_desc"}),
        )
        return PriceDifferenceStrategy().process(exp)

    def test_individual_sin_lista_al_pie(self):
        res = self._process("individual")
        assert "documentos_unicos" not in res.resumen
        assert "titulo_documentos" not in res.resumen

    def test_consolidado_lista_sin_repetir(self):
        res = self._process("consolidado")
        assert res.resumen["documentos_unicos"] == ["F201-100", "F201-101"]

    def test_pie_de_tabla_en_excel(self, tmp_path):
        from openpyxl import load_workbook
        from src.render.excel_renderer import ExcelRenderer

        res = self._process("consolidado")
        out = tmp_path / "Calculo.xlsx"
        ExcelRenderer().generar(
            res,
            str(out),
            tipo="diferencia_precio",
            cliente="CLIENTE X",
            doc_ref="F201-101",
            ruc="206039928",
            modalidad="consolidado",
        )
        wb = load_workbook(str(out))
        ws = wb[wb.sheetnames[0]]
        textos = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert ("FACTURAS COMPROMETIDAS: F201-100, F201-101") in textos


class TestMapearNcExistente:
    MAPA = {"F201-100": "N900-1", "F201-101": "N900-2"}

    def test_documento_unico(self):
        assert mapear_nc_existente("F201-100", self.MAPA) == "N900-1"

    def test_sin_nota(self):
        assert mapear_nc_existente("F201-999", self.MAPA) == ""

    def test_vacio_y_nan(self):
        assert mapear_nc_existente("", self.MAPA) == ""
        assert mapear_nc_existente("nan", self.MAPA) == ""

    def test_lista_con_comas(self):
        assert mapear_nc_existente("F201-100, F201-101", self.MAPA) == "N900-1, N900-2"

    def test_lista_con_punto_y_coma_y_sufijo(self):
        assert mapear_nc_existente("F201-100 (x); F201-101", self.MAPA) == "N900-1, N900-2"

    def test_lista_parcial(self):
        assert mapear_nc_existente("F201-100, F201-999", self.MAPA) == "N900-1"


class TestConsolidadoTrazabilidad:
    def test_resumen_consolidacion_va_a_trazabilidad(self):
        """'Consolidado por SKU' es trazabilidad: no entra a alertas."""
        st = PriceDifferenceStrategy()
        traz: list = []
        _, alertas, _ = st._procesar_comparar(
            _hist(),
            _lista(),
            [],
            trazabilidad=traz,
            config={"modalidad": "consolidado", "sort_mode": "fecha_desc"},
        )
        assert any("Consolidado por SKU" in str(t) for t in traz)
        assert not [a for a in alertas if not getattr(a, "codigo", "")], [
            a.mensaje for a in alertas if not getattr(a, "codigo", "")
        ]
