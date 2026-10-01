"""DO consolidado: misma convención de consolidado que DC y VRS.

- una fila por SKU con FACTURAS (lista) en vez de FACTURA;
- precio por MODA de los precios atendidos y detalle del corte en la auditoría
  (incluido el % de descuento aplicado);
- monto por suma exacta de las líneas (MONTO_EXACTO), sin lista de precios
  (PRECIO_BASE no existe en DO);
- las notas reconciliadas se detallan en la fila sin sumar AL12 al panel.
"""

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.core.nc_reconciliation import reconciliar_notas
from src.domain import RecognitionResult
from src.render.excel_renderer import ExcelRenderer
from src.strategies.price_difference import PriceDifferenceStrategy
from src.ui.config_builder import build_config


def _historial():
    return pd.DataFrame(
        [
            {
                "CODIGO": "S1",
                "ARTICULO": "Prod 1",
                "LINEA": "L01",
                "CANTIDAD": 50.0,
                "SOLES": 170.0,
                "FECHA": pd.Timestamp("2026-01-05"),
                "TIPO_DOC": "F",
                "SERIE": "F201",
                "NUMERO": "100",
                "DOC_ID": "FF201-100",
                "COD_CLIENTE": "1",
                "DOC_CLIENTE": "206039928",
                "TIPO_CLASE": "factura",
                "FACTURA_REF": "",
                "CANTIDAD_FAE": 0.0,
            },
            {
                "CODIGO": "S1",
                "ARTICULO": "Prod 1",
                "LINEA": "L01",
                "CANTIDAD": 30.0,
                "SOLES": 102.0,
                "FECHA": pd.Timestamp("2026-01-10"),
                "TIPO_DOC": "F",
                "SERIE": "F201",
                "NUMERO": "101",
                "DOC_ID": "FF201-101",
                "COD_CLIENTE": "1",
                "DOC_CLIENTE": "206039928",
                "TIPO_CLASE": "factura",
                "FACTURA_REF": "",
                "CANTIDAD_FAE": 0.0,
            },
            {
                "CODIGO": "S2",
                "ARTICULO": "Prod 2",
                "LINEA": "L02",
                "CANTIDAD": 10.0,
                "SOLES": 40.0,
                "FECHA": pd.Timestamp("2026-01-15"),
                "TIPO_DOC": "F",
                "SERIE": "F201",
                "NUMERO": "100",
                "DOC_ID": "FF201-100",
                "COD_CLIENTE": "1",
                "DOC_CLIENTE": "206039928",
                "TIPO_CLASE": "factura",
                "FACTURA_REF": "",
                "CANTIDAD_FAE": 0.0,
            },
        ]
    )


def _desc():
    return pd.DataFrame({"CODIGO": ["S1", "S2"], "DESCUENTO": [0.05, 0.10]})


def _nota_fae_exacta():
    return pd.DataFrame(
        [
            {
                "CODIGO": "S1",
                "ARTICULO": "Prod 1",
                "LINEA": "L01",
                "CANTIDAD": 0.0,
                "SOLES": -34.0,
                "FECHA": pd.Timestamp("2026-01-25"),
                "TIPO_DOC": "NC",
                "SERIE": "F900",
                "NUMERO": "1",
                "DOC_ID": "FF900-1",
                "TIPO_NOTA": "FAE",
                "CATEGORIA": "valor",
                "FACTURA_REF": "FF201-100",
                "DOC_NOTA": "FF900-1",
                "CANTIDAD_FAE": 50.0,
                "COD_CLIENTE": "1",
                "DOC_CLIENTE": "206039928",
                "TIPO_CLASE": "nc",
            }
        ]
    )


def _run(modalidad="consolidado", con_nota=False, global_pct=None):
    hist = _historial()
    documentos = {
        "facturas": {"mostrar": True, "usar": True},
        "nc": {"mostrar": True, "usar": True},
        "ndb": {"mostrar": True, "usar": False},
    }
    full = pd.concat([hist, _nota_fae_exacta()], ignore_index=True) if con_nota else hist
    rec = reconciliar_notas(full, documentos=documentos, modalidad=modalidad)
    config = {
        "modalidad": modalidad,
        "documentos_historial": documentos,
        "reconciliacion_nc": rec,
        "sort_mode": "fecha_desc",
    }
    alertas: list = []
    strat = PriceDifferenceStrategy()
    if global_pct is not None:
        df, row_alerts = strat._procesar_descuento_global(hist, global_pct, alertas, config=config)
    else:
        df, row_alerts = strat._procesar_descuento_simple(hist, _desc(), alertas, config=config)
    return df, row_alerts


def _xlsx(tmp_path, df, modalidad="consolidado"):
    res = RecognitionResult(
        dataframe=df,
        dataframe_excel=df.copy(),
        resumen={
            "total_nc": float(df["MONTO_NC"].sum()),
            "skus_afectados": df["SKU"].nunique(),
            "doc_ref": "FF201-100",
        },
    )
    out = tmp_path / f"DO_{modalidad}.xlsx"
    ExcelRenderer().generar(
        res,
        str(out),
        tipo="descuento_precio",
        cliente="CLIENTE X",
        doc_ref="FF201-100",
        modalidad=modalidad,
    )
    wb = load_workbook(str(out))
    ws = wb[wb.sheetnames[0]]
    grid = [[c.value for c in row] for row in ws.iter_rows()]
    head = next(i for i, r in enumerate(grid) if "% DESC." in r)
    return grid[head], grid[head + 1 :]


class TestDoConsolidado:
    def test_una_fila_por_sku(self):
        df, _ = _run()
        assert sorted(df["SKU"]) == ["S1", "S2"]
        assert df.set_index("SKU")["CANTIDAD"].to_dict() == {"S1": 80.0, "S2": 10.0}

    def test_precios_por_moda(self):
        hist = _historial()
        hist.loc[hist["DOC_ID"] == "FF201-101", "SOLES"] = 96.0  # 3.20/u
        # tercera línea de S1 al mismo precio de factura -> moda sin empate
        extra = hist[hist["DOC_ID"] == "FF201-100"].iloc[[0]].copy()
        extra["CANTIDAD"] = 20.0
        extra["SOLES"] = 68.0
        hist = pd.concat([hist, extra], ignore_index=True)
        config = {"modalidad": "consolidado", "sort_mode": "fecha_desc"}
        df, _ = PriceDifferenceStrategy()._procesar_descuento_simple(
            hist, _desc(), [], config=config
        )
        fila = df[df["SKU"] == "S1"].iloc[0]
        assert fila["PRECIO_HIST"] == 3.40  # moda (3.40 x 2, 3.20 x 1)
        assert "rango S/ 3.20000" in fila["AUDITORIA_NC"]
        # el monto manda: suma exacta de las líneas, no moda × cantidad
        # (50u a 3.40 → 8.50) + (30u a 3.20 → 4.80) + (20u a 3.40 → 3.40)
        assert fila["MONTO_NC"] == pytest.approx(8.50 + 4.80 + 3.40, abs=0.01)
        # ... y no 100u × (3.40 - 3.23) = 17.00 que saldría de la moda
        assert fila["DIFERENCIA"] == pytest.approx(0.17, abs=0.001)
        assert fila["MONTO_EXACTO"] == fila["MONTO_NC"]

    def test_auditoria_detalla_facturas_moda_y_descuento(self):
        df, _ = _run()
        audit = df[df["SKU"] == "S1"].iloc[0]["AUDITORIA_NC"]
        assert "Consolidado 2 líneas de 2 factura(s)" in audit
        assert "FF201-100, FF201-101" in audit
        assert "moda S/ 3.40000" in audit
        assert "desc. aplicado 5.00%" in audit
        assert "Cortes:" in audit

    def test_facturas_lista_y_factura_principal(self):
        df, _ = _run()
        fila = df[df["SKU"] == "S1"].iloc[0]
        assert fila["FACTURAS"] == "FF201-100, FF201-101"
        assert fila["FACTURA"] == "FF201-100"  # la de mayor SOLES
        assert fila["LINEA"] == "L01"

    def test_porcentaje_de_descuento_viaja_a_la_fila(self):
        df, _ = _run()
        assert df.set_index("SKU")["%_DESCUENTO"].to_dict() == {"S1": 0.05, "S2": 0.10}

    def test_no_requiere_precio_base(self):
        """DO no tiene lista de precios: consolidar no debe inventar columnas."""
        df, _ = _run()
        assert "PRECIO_BASE" not in df.columns
        assert "DESCUENTO_COMPUESTO" not in df.columns
        assert (df["MONTO_NC"] > 0).all()

    def test_modo_descuento_global_consolida_igual(self):
        df, _ = _run(global_pct=0.05)
        assert sorted(df["SKU"]) == ["S1", "S2"]
        fila = df[df["SKU"] == "S1"].iloc[0]
        assert fila["%_DESCUENTO"] == 0.05
        assert fila["MONTO_NC"] == pytest.approx(80 * 0.17, abs=0.01)
        assert "desc. aplicado 5.00%" in fila["AUDITORIA_NC"]

    def test_nota_exacta_se_detalla_sin_alerta_de_panel(self):
        df, alertas = _run(modalidad="consolidado", con_nota=True)
        fila = df[df["SKU"] == "S1"].iloc[0]
        assert "no aplicado" in fila["ALERTA"]
        assert "NF900-1" in fila["AUDITORIA_NC"]
        # consolidado: la nota no ajusta (precio de factura) ni avisa en el panel
        assert fila["PRECIO_HIST"] == 3.40
        assert not [a for a in alertas if a.codigo == "AL12"]


class TestDoConsolidadoExcel:
    def test_columna_facturas_y_monto_estatico(self, tmp_path):
        df, _ = _run()
        headers, filas = _xlsx(tmp_path, df)
        assert "FACTURAS" in headers and "FACTURA" not in headers
        por_sku = {r[headers.index("SKU")]: r for r in filas if r[headers.index("SKU")]}
        fila_s1 = por_sku["S1"]
        assert fila_s1[headers.index("FACTURAS")] == "FF201-100, FF201-101"
        # % DESC. con valor (no 0) y MONTO exacto escrito como número
        assert fila_s1[headers.index("% DESC.")] == 0.05
        monto = float(df.set_index("SKU")["MONTO_NC"]["S1"])
        assert fila_s1[headers.index("MONTO")] == pytest.approx(monto, abs=0.01)
        assert not str(fila_s1[headers.index("MONTO")]).startswith("=")

    def test_porcentaje_descuento_se_escribe_en_individual(self, tmp_path):
        df, _ = _run(modalidad="individual")
        headers, filas = _xlsx(tmp_path, df, modalidad="individual")
        assert "% DESC." in headers
        assert filas[0][headers.index("% DESC.")] == 0.05
        # en individual MONTO sigue siendo la fórmula viva del motor
        assert str(filas[0][headers.index("MONTO")]).startswith("=ROUND(")


class TestDoConfig:
    def test_modalidad_llega_desde_la_ui(self):
        for modalidad in ("individual", "consolidado"):
            cfg = build_config("descuento_precio", {"modalidad": modalidad})
            assert cfg["modalidad"] == modalidad


class TestDoPorcentajeDescEnExcel:
    """Regresión: `generar()` renombra %_DESCUENTO → DESC1 y la celda % DESC.

    quedaba en 0, así que las fórmulas de PRECIO NETO / DIF. UNITARIA / MONTO
    de la hoja se resolvían contra 0 (NC en 0) aunque el motor calculaba bien.
    """

    def test_descuento_factura_tambien_lo_escribe(self, tmp_path):
        """El mismo rename afectaba a descuento_factura (% DESC. del archivo)."""
        df = pd.DataFrame(
            {
                "SKU": ["S1"],
                "ARTICULO": ["Prod 1"],
                "CANTIDAD": [100],
                "FACTURA": ["FF201-100"],
                "ALERTA": ["OK"],
                "PRECIO_UNITARIO": [3.40],
                "%_DESCUENTO": [0.05],
                "MONTO_NC": [17.00],
            }
        )
        res = RecognitionResult(
            dataframe=df,
            dataframe_excel=df.copy(),
            resumen={"total_nc": 17.0, "skus_afectados": 1, "doc_ref": "FF201-100"},
        )
        out = tmp_path / "df.xlsx"
        ExcelRenderer().generar(
            res, str(out), tipo="descuento_factura", cliente="CLIENTE X", doc_ref="FF201-100"
        )
        wb = load_workbook(str(out))
        ws = wb[wb.sheetnames[0]]
        grid = [[c.value for c in row] for row in ws.iter_rows()]
        head = next(i for i, r in enumerate(grid) if "% DESC." in r)
        headers = grid[head]
        assert grid[head + 1][headers.index("% DESC.")] == 0.05


class TestDoDocumentosYAlertas:
    """documentos_unicos solo facturas (P2) y alerta de precios con AL03."""

    def test_documentos_unicos_solo_facturas(self):
        """La nota reconciliada ajusta cantidades pero no entra en FACTURAS."""
        merged = pd.DataFrame(
            [
                {
                    "CODIGO": "S1",
                    "ARTICULO": "Prod 1",
                    "LINEA": "L01",
                    "FECHA": pd.Timestamp("2026-01-05"),
                    "CANTIDAD": 100.0,
                    "SOLES": 1000.0,
                    "PRECIO_HIST": 10.0,
                    "PRECIO_NETO": 9.5,
                    "MONTO_NC": 50.0,
                    "FACTURA": "F204-100",
                    "ALERTA": "",
                    "ALERTA_NOTAS": "",
                    "AUDITORIA_NOTAS": "",
                    "TIPO_CLASE": "factura",
                },
                {
                    "CODIGO": "S1",
                    "ARTICULO": "Prod 1",
                    "LINEA": "L01",
                    "FECHA": pd.Timestamp("2026-01-20"),
                    "CANTIDAD": -10.0,
                    "SOLES": -100.0,
                    "PRECIO_HIST": 10.0,
                    "PRECIO_NETO": 10.0,
                    "MONTO_NC": 0.0,
                    "FACTURA": "NN204-900010",
                    "ALERTA": "",
                    "ALERTA_NOTAS": "",
                    "AUDITORIA_NOTAS": "",
                    "TIPO_CLASE": "nc",
                },
            ]
        )
        df = PriceDifferenceStrategy()._consolidar_por_sku(merged, sort_mode="fecha_desc")
        row = df.iloc[0]
        assert row["FACTURAS"] == "F204-100"
        assert row["FACTURA"] == "F204-100"

    def test_alerta_precios_distintos_tiene_al03(self):
        """Precios netos distintos en consolidado = AL03 (no codigo vacio)."""
        documentos = {
            "facturas": {"mostrar": True, "usar": True},
            "nc": {"mostrar": True, "usar": True},
            "ndb": {"mostrar": True, "usar": False},
        }
        hist = pd.DataFrame(
            [
                {
                    "CODIGO": "S1",
                    "ARTICULO": "Prod 1",
                    "LINEA": "L01",
                    "CANTIDAD": 50.0,
                    "SOLES": 200.0,
                    "FECHA": pd.Timestamp("2026-01-05"),
                    "TIPO_DOC": "F",
                    "SERIE": "F201",
                    "NUMERO": "100",
                    "DOC_ID": "FF201-100",
                    "COD_CLIENTE": "1",
                    "DOC_CLIENTE": "206039928",
                    "TIPO_CLASE": "factura",
                    "FACTURA_REF": "",
                    "CANTIDAD_FAE": 0.0,
                },
                {
                    "CODIGO": "S1",
                    "ARTICULO": "Prod 1",
                    "LINEA": "L01",
                    "CANTIDAD": 30.0,
                    "SOLES": 150.0,
                    "FECHA": pd.Timestamp("2026-01-10"),
                    "TIPO_DOC": "F",
                    "SERIE": "F201",
                    "NUMERO": "101",
                    "DOC_ID": "FF201-101",
                    "COD_CLIENTE": "1",
                    "DOC_CLIENTE": "206039928",
                    "TIPO_CLASE": "factura",
                    "FACTURA_REF": "",
                    "CANTIDAD_FAE": 0.0,
                },
            ]
        )
        rec = reconciliar_notas(hist, documentos=documentos, modalidad="consolidado")
        config = {
            "modalidad": "consolidado",
            "documentos_historial": documentos,
            "reconciliacion_nc": rec,
            "sort_mode": "fecha_desc",
        }
        alertas: list = []
        _, row_alerts = PriceDifferenceStrategy()._procesar_descuento_simple(
            hist, _desc(), alertas, config=config
        )
        al03 = [a for a in row_alerts if getattr(a, "codigo", "") == "AL03"]
        assert al03, [str(a.mensaje)[:60] for a in row_alerts]
        assert all(getattr(a, "codigo", "") for a in row_alerts if a.tipo == "warning")
