"""Histórico del expediente: una hoja por año cuando abarca varios.

Con más de 12 meses el archivo se volvía pesado de abrir (una sola hoja con
decenas de miles de filas). Ahora se parte en una hoja por año calendar.

Lo que estos tests fijan, y que ya se rompió una vez al implementar:

- Con un solo año se conserva la hoja única "Historico" (el caso habitual no
  cambia).
- Cada hoja recalcula sus propios agregados: RESUMEN y TOTAL son de ESA hoja.
- El sombreado gris de las NC no se desalinea. `nota_flags` es una lista
  posicional: si se parte el dataframe sin recalcularla, el gris cae en las
  filas equivocadas y no da ningún error.
- Las filas sin fecha interpretable no se pierden ni se reparten.
- El nombre de hoja se sanea (Excel prohíbe []:*?/\\ y tops de 31 chars).
"""

from __future__ import annotations

from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.core.fechas import excel_fmt_ui
from src.render.audit_renderer import _agrupar_por_anio, _nombre_seguro, generar_historico_clasico

COLS = (
    "FECHA",
    "TIPO_DOC",
    "SERIE",
    "NUMERO",
    "CODIGO",
    "ARTICULO",
    "CANTIDAD",
    "PRECIO_UNITARIO",
    "SOLES",
    "DOC_CLIENTE",
)


def _df(filas):
    return pd.DataFrame([dict(zip(COLS, f)) for f in filas])


def _fila(fecha, sku, cantidad, soles, tipo="F001", serie="A001", numero="1001"):
    return (fecha, tipo, serie, numero, sku, f"ART {sku}", cantidad, 5.0, soles, "20100047218")


MULTI = _df(
    [
        _fila("15/01/2024", "SKU1", 10, 50.0),
        _fila("20/03/2024", "SKU2", 5, 20.0, numero="1002"),
        _fila("05/06/2024", "SKU3", -3, -18.0, tipo="NC", serie="B001", numero="9001"),
        _fila("02/02/2025", "SKU4", 8, 60.0, numero="2001"),
        _fila("11/11/2025", "SKU5", 4, 32.0, numero="2002"),
        _fila("07/03/2026", "SKU6", 12, 42.0, tipo="F003", numero="3001"),
        _fila("09/09/2026", "SKU7", -2, -18.0, tipo="NDB", serie="B001", numero="9002"),
    ]
)


def _generar(df, tmp_path, **kw):
    salida = tmp_path / "Historico.xlsx"
    generar_historico_clasico(
        df,
        cliente_nombre="DEMO S.A.",
        cliente_ruc="20100047218",
        tipo_operacion="01",
        fecha_desde="01-01-2024",
        fecha_hasta="30-09-2026",
        vendedor="01188",
        expediente_id="EXP-TEST",
        ruta_salida=salida,
        **kw,
    )
    return load_workbook(salida)


def _registros_de(ws) -> int:
    """Cantidad de filas de datos declarada en el RESUMEN de la hoja."""
    texto = ws.cell(row=11, column=1).value or ""
    return int(texto.split("registro(s)")[0].split(":")[-1].strip())


# ── Partición ────────────────────────────────────────────────────────────────


def test_un_solo_anio_conserva_la_hoja_historico(tmp_path):
    """El caso habitual no cambia: una hoja, llamada 'Historico'."""
    df = _df([_fila("15/01/2025", "S1", 1, 10.0), _fila("20/03/2025", "S2", 1, 20.0)])
    wb = _generar(df, tmp_path)
    assert wb.sheetnames == ["Historico"]


def test_varios_anos_una_hoja_por_anio(tmp_path):
    wb = _generar(MULTI, tmp_path)
    assert wb.sheetnames == ["2024", "2025", "2026"]


def test_orden_de_hojas_es_cronologico(tmp_path):
    assert _generar(MULTI, tmp_path).sheetnames == sorted(_generar(MULTI, tmp_path).sheetnames)


# ── Agregados por hoja ───────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "hoja,registros,facturas,notas,soles",
    [("2024", 3, 2, 1, 52.0), ("2025", 2, 2, 0, 92.0), ("2026", 2, 1, 1, 24.0)],
)
def test_resumen_y_totales_son_de_cada_hoja(tmp_path, hoja, registros, facturas, notas, soles):
    """Si los totales se calcularan sobre el dataframe entero, cada hoja
    repetiría el total global y el RESUMEN mentiría."""
    wb = _generar(MULTI, tmp_path)
    ws = wb[hoja]
    resumen = ws.cell(row=11, column=1).value or ""
    assert f"{registros} registro(s)" in resumen
    assert f"{facturas} factura(s)" in resumen
    assert f"{notas} nota(s)" in resumen
    assert f"{soles:,.2f}" in resumen

    # TOTAL va en 13 + n_filas (los datos arrancan en 13).
    total = ws.cell(row=13 + registros, column=9).value
    assert abs(total - soles) < 0.01


def test_el_sombreado_de_notas_no_se_desalinea(tmp_path):
    """`nota_flags` es posicional: al partir hay que recalcularla por hoja.

    Es el fallo más peligroso de este refactor porque no lanza error: el gris
    cae en filas equivocadas y el archivo parece correcto.
    """
    wb = _generar(MULTI, tmp_path)
    esperados = {"2024": ["SKU3"], "2025": [], "2026": ["SKU7"]}
    for hoja, esperadas in esperados.items():
        ws = wb[hoja]
        n = _registros_de(ws)
        grises = []
        for r in range(13, 13 + n):
            rgb = getattr(getattr(ws.cell(row=r, column=1).fill, "fgColor", None), "rgb", None)
            if rgb and "E7E6E6" in str(rgb).upper():
                grises.append(ws.cell(row=r, column=5).value)
        assert grises == esperadas, f"hoja {hoja}: {grises} != {esperadas}"


def test_el_resaltado_de_reclamo_queda_en_su_hoja(tmp_path):
    """La clave del reclamo es (factura, SKU); debe caer en su año."""
    wb = _generar(MULTI, tmp_path, reclamados={("FA001-1001", "SKU1")})
    ws = wb["2024"]
    amarillas = []
    for r in range(13, 16):
        rgb = getattr(getattr(ws.cell(row=r, column=1).fill, "fgColor", None), "rgb", None)
        if rgb and "FFF2CC" in str(rgb).upper():
            amarillas.append(ws.cell(row=r, column=5).value)
    assert amarillas == ["SKU1"]
    assert "SKU4" not in amarillas, "el reclamo se pintó en la hoja equivocada"


def test_cabecera_fija_y_area_de_impresion_por_hoja(tmp_path):
    wb = _generar(MULTI, tmp_path)
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        assert ws.freeze_panes == "A13", hoja
        assert ws.print_area.endswith("$J$%d" % int(ws.print_area.split("$J$")[1]))


# ── Casos borde ──────────────────────────────────────────────────────────────


def test_filas_sin_fecha_no_se_pierden():
    df = _df([_fila("15/01/2024", "S1", 1, 10.0), _fila("basura", "S2", 1, 20.0)])
    grupos = _agrupar_por_anio(df)
    nombres = [g[0] for g in grupos]
    assert "2024" in nombres
    assert "Sin fecha" in nombres
    assert sum(len(g[1]) for g in grupos) == 2


def test_dataframe_vacio_no_rompe(tmp_path):
    wb = _generar(pd.DataFrame(), tmp_path)
    assert len(wb.sheetnames) == 1


def test_sin_columna_fecha_no_rompe():
    df = _df([_fila("15/01/2024", "S1", 1, 10.0)]).drop(columns=["FECHA"])
    grupos = _agrupar_por_anio(df)
    assert [g[0] for g in grupos] == ["Historico"]


def test_nombre_de_hoja_se_sanea():
    usados = set()
    assert "/" not in _nombre_seguro("2024/2025", usados)
    assert len(_nombre_seguro("x" * 50, usados)) <= 31
    a = _nombre_seguro("2024", usados)
    b = _nombre_seguro("2024", usados)
    assert a != b, "dos hojas no pueden llamarse igual"


def test_acepta_fechas_con_hora():
    """Varias columnas del ERP traen la hora pegada."""
    df = _df(
        [
            _fila("15/01/2024 08:30:00", "S1", 1, 10.0),
            _fila("2024-06-01T10:00:00Z", "S2", 1, 20.0),
            _fila("03/07/2025 23:59:00", "S3", 1, 30.0),
        ]
    )
    nombres = [g[0] for g in _agrupar_por_anio(df)]
    assert nombres == ["2024", "2025"]


# ── Formato de fecha visible ────────────────────────────────────────────────


@pytest.mark.parametrize(
    "origen",
    ["15/01/2024", "2024-01-15", "15-01-2024", "15/01/2024 08:30:00", datetime(2024, 1, 15)],
)
def test_la_columna_de_fecha_se_normaliza_a_dd_mm_yyyy(tmp_path, origen):
    """La columna FECHA EMISION sale siempre segun FMT_UI.

    Antes pasaba la cadena tal cual, así que el mismo histórico mezclaba barras,
    ISO y guiones según cómo lo guardó el ERP. Es lo mismo que el resto de la
    app (src/core/fechas.FMT_UI).
    """
    from src.core.fechas import fecha_ui

    df = _df([_fila(origen, "S1", 1, 10.0)])
    ws = _generar(df, tmp_path).active
    celda = ws.cell(row=13, column=1)
    # Celda de fecha REAL (no texto): asi Excel la ordena y filtra.
    assert hasattr(celda.value, "year"), celda.value
    assert celda.number_format == excel_fmt_ui()
    assert fecha_ui(celda.value) == "15/01/2024"


def test_una_fecha_ilegible_no_se_pierde(tmp_path):
    """Si no se puede interpretar, se conserva el texto: es preferible una
    fecha rara a perder el dato."""
    df = _df([_fila("sin fecha", "S1", 1, 10.0)])
    ws = _generar(df, tmp_path).active
    assert ws.cell(row=13, column=1).value == "sin fecha"
