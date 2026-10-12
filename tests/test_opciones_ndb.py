"""Opciones de calculo NDB: solo-manual y diferencia inversa.

1. Solo-NDB es MANUAL. Antes habia un automatico (`nc_total <= 0.001`) que
   sacaba unicamente CalculoND sin avisar: el usuario esperaba el expediente
   normal y recibia otro. Ahora sin el check nunca se fuerza ese modo.

2. El CalculoND muestra la DIFERENCIA INVERSA por unidad (NETO - FACTURADO,
   en positivo). En NC la diferencia es HIST-NETO (lo cobrado de mas); en
   NDB es al reves (lo cobrado de menos). Un debito es un importe a cobrar,
   asi que se muestra en positivo, no como un negativo.
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.ui.expediente_service import _decidir_solo_nd, _generar_calculo_nd


def _nd_df():
    return pd.DataFrame(
        [
            {
                "FACTURA": "F201-1",
                "SKU": "S1",
                "ARTICULO": "Prod 1",
                "CANTIDAD": 10,
                "PRECIO_HIST": 2.00,
                "PRECIO_NETO": 2.50,
                "DIFERENCIA": -0.50,
                "MONTO_ND": 5.00,
            },
            {
                "FACTURA": "F201-2",
                "SKU": "S2",
                "ARTICULO": "Prod 2",
                "CANTIDAD": 4,
                "PRECIO_HIST": 1.00,
                "PRECIO_NETO": 1.75,
                "DIFERENCIA": -0.75,
                "MONTO_ND": 3.00,
            },
        ]
    )


def _hoja(salida: Path):
    wb = load_workbook(salida, data_only=True)
    ws = wb.active
    fila_h = next(r for r in range(1, 20) if ws.cell(row=r, column=1).value == "N°")
    return ws, fila_h


class TestSoloNdEsManual:
    def test_sin_check_no_hay_solo_nd_aunque_no_haya_nc(self):
        """El automatico viejo: nc_total en cero forzaba el modo sin avisar."""
        assert _decidir_solo_nd(solo_ndb=False, hay_nd=True) is False

    def test_con_check_y_negativas_si(self):
        assert _decidir_solo_nd(solo_ndb=True, hay_nd=True) is True

    def test_con_check_pero_sin_negativas_no(self):
        """Marcarlo sin negativas que calcular no tiene sentido."""
        assert _decidir_solo_nd(solo_ndb=True, hay_nd=False) is False

    def test_sin_nada_no(self):
        assert _decidir_solo_nd(solo_ndb=False, hay_nd=False) is False


class TestNegativasSinNdb:
    """Con ocultar 0.00 y sin NDB, el expediente NC tiene que ser todo positivo.

    Caso real: 78307 con DIFERENCIA -0.02357 (2.29080 facturado vs 2.31437
    neto) salia con AL02 "Diferencia negativa, sin NC" en un expediente que
    debia ser todo positivo. Sin NDB esa fila no es candidata a nada: es
    ruido. Con NDB se conserva porque alimenta el CalculoND.
    """

    def _exp(self):
        from types import SimpleNamespace

        import pandas as pd

        from src.domain import RecognitionResult

        df = pd.DataFrame(
            [
                {
                    "CODIGO": "78307",
                    "CLIENTE": "C1",
                    "DIFERENCIA": -0.02357,
                    "MONTO_NC": 0.0,
                    "ALERTA": "AL02 - Diferencia negativa, sin NC",
                },
                {
                    "CODIGO": "OTRO",
                    "CLIENTE": "C1",
                    "DIFERENCIA": 1.50,
                    "MONTO_NC": 150.0,
                    "ALERTA": "AL01",
                },
            ]
        )
        res = RecognitionResult(
            dataframe=df,
            resumen={"total_nc": 150.0},
            alertas=[],
            trazabilidad=[],
        )
        return SimpleNamespace(resultado=res, datos={}, contexto=None)

    def test_sin_ndb_la_negativa_no_sale(self):
        from src.ui.expediente_service import _clonar_sin_cero

        clon = _clonar_sin_cero(self._exp(), conservar_negativas=False)
        df = clon.resultado.dataframe
        assert "78307" not in df["CODIGO"].tolist()
        assert bool((df["DIFERENCIA"] > 0).all()), "tiene que ser todo positivo"

    def test_con_ndb_la_negativa_se_conserva(self):
        from src.ui.expediente_service import _clonar_sin_cero

        clon = _clonar_sin_cero(self._exp(), conservar_negativas=True)
        df = clon.resultado.dataframe
        assert "78307" in df["CODIGO"].tolist()
        assert "OTRO" in df["CODIGO"].tolist()

    def test_por_defecto_conserva_para_no_romper_llamados_viejos(self):
        """El default True preserva el comportamiento anterior."""
        from src.ui.expediente_service import _clonar_sin_cero

        clon = _clonar_sin_cero(self._exp())
        assert "78307" in clon.resultado.dataframe["CODIGO"].tolist()

    def test_la_hoja_tiene_dif_unitaria(self, tmp_path):
        salida = tmp_path / "CalculoND.xlsx"
        _generar_calculo_nd(_nd_df(), "diferencia_precio", "CLI", "V", "F201-1", "RUC", salida)
        ws, fila_h = _hoja(salida)
        headers = [ws.cell(row=fila_h, column=c).value for c in range(1, 10)]
        assert "DIF. UNITARIA" in headers
        assert "MONTO ND" in headers
        assert len([h for h in headers if h]) == 9

    def test_la_diferencia_es_positiva_y_consistente(self, tmp_path):
        salida = tmp_path / "CalculoND.xlsx"
        _generar_calculo_nd(_nd_df(), "diferencia_precio", "CLI", "V", "F201-1", "RUC", salida)
        ws, fila_h = _hoja(salida)
        for r in (fila_h + 1, fila_h + 2):
            hist = ws.cell(row=r, column=6).value
            neto = ws.cell(row=r, column=7).value
            dif = ws.cell(row=r, column=8).value
            monto = ws.cell(row=r, column=9).value
            cant = ws.cell(row=r, column=5).value
            assert dif is not None and dif > 0, "la inversa se muestra en positivo"
            assert abs(dif - round(neto - hist, 6)) < 1e-6
            assert abs(monto - round(dif * cant, 2)) < 0.01

    def test_el_total_suma_los_montos(self, tmp_path):
        salida = tmp_path / "CalculoND.xlsx"
        _generar_calculo_nd(_nd_df(), "diferencia_precio", "CLI", "V", "F201-1", "RUC", salida)
        ws, fila_h = _hoja(salida)
        fila_t = next(
            r for r in range(fila_h, 25) if ws.cell(row=r, column=1).value == "TOTAL ND (S/)"
        )
        assert abs(ws.cell(row=fila_t, column=9).value - 8.00) < 0.01

    def test_el_titulo_cubre_las_9_columnas(self, tmp_path):
        salida = tmp_path / "CalculoND.xlsx"
        _generar_calculo_nd(_nd_df(), "diferencia_precio", "CLI", "V", "F201-1", "RUC", salida)
        ws, _ = _hoja(salida)
        merges = [str(m) for m in ws.merged_cells.ranges]
        assert "A1:I1" in merges, "el titulo tiene que cubrir las 9 columnas"
class TestNcYNdbIndependientes:
    """NC y NDB son independientes: las negativas nunca entran a los
    archivos normales, haya o no NDB pedido.

    Antes, con NDB prendido, `_clonar_sin_cero(conservar_negativas=True)`
    dejaba las negativas en `df_res` y estas entraban al Informe/Calculo/
    Historico. Ahora van unicamente al CalculoND via `df_nd`.
    """

    def _exp_mixto(self):
        from types import SimpleNamespace

        import pandas as pd

        from src.domain import RecognitionResult

        df = pd.DataFrame(
            [
                {
                    "CODIGO": "POS",
                    "CLIENTE": "C1",
                    "DIFERENCIA": 2.0,
                    "MONTO_NC": 20.0,
                    "ALERTA": "AL01",
                },
                {
                    "CODIGO": "NEG",
                    "CLIENTE": "C1",
                    "DIFERENCIA": -1.0,
                    "MONTO_NC": 0.0,
                    "ALERTA": "AL02 - Diferencia negativa, sin NC",
                },
            ]
        )
        res = RecognitionResult(
            dataframe=df, resumen={"total_nc": 20.0}, alertas=[], trazabilidad={}
        )
        return SimpleNamespace(resultado=res, datos={}, contexto=None)

    def test_los_archivos_normales_no_llevan_negativas(self):
        from src.ui.expediente_service import _clonar_solo_positivas

        clon = _clonar_solo_positivas(self._exp_mixto())
        df = clon.resultado.dataframe
        assert "POS" in df["CODIGO"].tolist()
        assert "NEG" not in df["CODIGO"].tolist()
        assert bool((df["DIFERENCIA"] >= 0).all())

    def test_sin_positivas_no_hay_archivos_normales(self):
        from src.ui.expediente_service import _clonar_solo_positivas

        from types import SimpleNamespace

        import pandas as pd

        from src.domain import RecognitionResult

        df = pd.DataFrame(
            [
                {
                    "CODIGO": "NEG",
                    "CLIENTE": "C1",
                    "DIFERENCIA": -1.0,
                    "MONTO_NC": 0.0,
                    "ALERTA": "AL02",
                }
            ]
        )
        res = RecognitionResult(
            dataframe=df, resumen={"total_nc": 0.0}, alertas=[], trazabilidad={}
        )
        exp = SimpleNamespace(resultado=res, datos={}, contexto=None)
        assert _clonar_solo_positivas(exp) is None

    def test_solo_nd_sin_positivas_genera_el_calculond(self, tmp_path):
        from src.ui.expediente_service import _solo_nd_sin_positivas

        import pandas as pd

        df_nd = pd.DataFrame(
            [
                {
                    "CLIENTE": "C1",
                    "FACTURA": "F201-1",
                    "SKU": "S1",
                    "ARTICULO": "P1",
                    "CANTIDAD": 10,
                    "PRECIO_HIST": 2.0,
                    "PRECIO_NETO": 2.5,
                    "DIFERENCIA": -0.5,
                    "MONTO_ND": 5.0,
                },
            ]
        )
        dirs = _solo_nd_sin_positivas(
            df_nd=df_nd,
            tipo_actual="diferencia_precio",
            df_historial=None,
            cliente_value="C1",
            vendedor_value="V",
            vendedor_display="V",
            desktop_path=tmp_path,
        )
        assert len(dirs) == 1
        nombres = [p.name for p in dirs[0].iterdir() if p.is_file()]
        assert len(nombres) == 1 and "CalculoND" in nombres[0]
        assert not any("Informe" in n or n.startswith("EXP-") and "Calculo." in n for n in nombres)

